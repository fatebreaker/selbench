"""Source data for every panel of the main figures (Figs 1-5), the Extended Data figures (ED Figs 1-2) and the
supplementary figures (Figs S1-S8).

The figures are not re-implemented here. The figure list and captions come from manuscript/sections/figures.tex,
manuscript/sections/extended_data.tex and manuscript/supplement_figures.tex; the function that saves each figure file
is found in 64_figures_revision.py (main and Extended Data figures, the worked example) or 13_figures.py (the other
supplementary figures). Each is run once with matplotlib instrumented:

* every Axes plotting call made by the figure code (scatter, errorbar, plot, bar/barh, lines, spans, boxplot, text,
  annotate, imshow, patches, tick labels, ...) is recorded with its arguments and a snapshot of the calling frame's
  local variables (the loop variables that say which model, dataset or class a mark belongs to);
* panel letters come from the module's panel_label(); unlabelled axes inherit the letter of the preceding labelled
  axes (Fig. 5a's six sub-axes) and inset axes that of their parent;
* every file read through pandas.read_csv / read_parquet / Path.read_text while a panel is drawn is recorded, which
  gives each panel's upstream files;
* save() is redirected to an in-memory PNG whose SHA-256 is compared with the published PNG in figures/ (byte
  identity shows that the exported values are those of the published figure); nothing is written to figures/.
  The module RNG is reset to its import-time state before each figure, as when a figure is drawn by its own
  invocation of the figure script (this fixes the drug-bootstrap intervals drawn in Fig. 2b,c).

Per-panel converters below (keyed by the drawing function) turn the recorded marks into tidy tables, naming columns
from the panel code (legend labels, tick labels, colours, loop variables). A panel without a converter falls back to
a generic dump of its marks, so a new panel is still exported. Plotted values are exported exactly as passed to
matplotlib; interval bounds come from the source row the panel code drew them from.

Outputs (manuscript/source_data/): Source_Data_Fig1.xlsx ... Source_Data_Fig5.xlsx, Source_Data_ExtendedDataFig1.xlsx,
Source_Data_ExtendedDataFig2.xlsx and Source_Data_SuppFig1.xlsx ... Source_Data_SuppFig8.xlsx (README sheet first,
then one sheet per panel; workbooks of figures that no longer exist are removed), the same tables as CSV under csv/,
README.csv, manifest.csv (file, panel, rows, sha256) and validation.csv (exported values against the upstream files;
byte identity of the regenerated and published figures; figures/ unchanged by this run).

Usage (from scripts/):  ../../.venv/bin/python 99i_source_data.py      [--dump: print the recorded marks only]
"""
import os

for _v in ('OMP_NUM_THREADS', 'OPENBLAS_NUM_THREADS', 'MKL_NUM_THREADS', 'NUMEXPR_NUM_THREADS'):
    os.environ.setdefault(_v, '1')

import copy
import datetime
import functools
import hashlib
import importlib
import inspect
import io
import json
import pathlib
import re
import sys
import warnings
from collections import OrderedDict
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib

matplotlib.use('Agg')
import matplotlib.image as mpimg
import matplotlib.pyplot as plt
from matplotlib.axes import Axes
from matplotlib.colors import to_hex
from matplotlib.lines import Line2D

warnings.filterwarnings('ignore', category=pd.errors.DtypeWarning)

SCRIPTS = Path(__file__).resolve().parent
PROJECT = SCRIPTS.parent
MANUSCRIPT = PROJECT / 'manuscript'
OUT = MANUSCRIPT / 'source_data'
FIGDIR = PROJECT / 'figures'
sys.path.insert(0, str(SCRIPTS))
sys.dont_write_bytecode = True        # write nothing outside manuscript/source_data/

M = importlib.import_module('64_figures_revision')
F = M.F
RNG_STATE0 = copy.deepcopy(M.RNG.bit_generator.state)       # the module RNG as at import (fresh seed)
MODULE_FILES = {str(Path(M.__file__).resolve()), str(Path(F.__file__).resolve())}
READ_CSV, READ_PARQUET, READ_TEXT = pd.read_csv, pd.read_parquet, pathlib.Path.read_text   # unpatched


def rel(p):
    p = Path(p).resolve()
    try:
        return str(p.relative_to(PROJECT))
    except ValueError:
        return str(p)


# =============================================================================================== recording
class Call:
    __slots__ = ('seq', 'ax', 'method', 'args', 'kw', 'func', 'lineno', 'ctx', 'stack', 'token')

    def __init__(self, **k):
        for a, v in k.items():
            setattr(self, a, v)

    def __repr__(self):
        return f'<{self.method} {self.func}:{self.lineno}>'

    @property
    def pfunc(self):
        """The drawing function: the outermost panel_* function on the stack, else the calling function."""
        p = [s for s in self.stack if s.startswith('panel_')]
        return p[-1] if p else self.func


class Recorder:
    def __init__(self):
        self.active = False
        self.suppress = 0
        self.ntok = 0
        self.token = None
        self.reset(None)

    def reset(self, key):
        self.fig = key
        self.calls = []
        self.labels = OrderedDict()   # id(ax) -> letter
        self.label_axes = {}          # id(ax) -> ax (kept alive)
        self.parents = {}             # id(inset) -> parent ax
        self.reads = []
        self.last_label = None
        self.saved = []               # (name, axes, png bytes)


REC = Recorder()

PLOT_METHODS = ['scatter', 'errorbar', 'plot', 'bar', 'barh', 'hlines', 'vlines', 'fill_between', 'fill_betweenx',
                'axvspan', 'axhspan', 'axvline', 'axhline', 'boxplot', 'violinplot', 'hist', 'step', 'stem', 'imshow',
                'pcolormesh', 'matshow', 'contour', 'contourf', 'text', 'annotate', 'add_patch', 'legend',
                'set_xticks', 'set_yticks', 'set_xticklabels', 'set_yticklabels', 'set_title', 'set_xlabel',
                'set_ylabel', 'invert_yaxis', 'invert_xaxis']
DATA_METHODS = {'scatter', 'errorbar', 'plot', 'bar', 'barh', 'hlines', 'vlines', 'fill_between', 'fill_betweenx',
                'axvspan', 'axhspan', 'axvline', 'axhline', 'boxplot', 'violinplot', 'hist', 'step', 'stem', 'imshow',
                'pcolormesh', 'matshow', 'contour', 'contourf', 'add_patch'}


def _stack_funcs(fr):
    out = []
    while fr is not None:
        if fr.f_code.co_filename in MODULE_FILES:
            out.append(fr.f_code.co_name)
        fr = fr.f_back
    return out


def _wrap_axes_method(name):
    orig = getattr(Axes, name)

    @functools.wraps(orig)
    def wrapper(self, *args, **kw):
        if REC.active and not REC.suppress:
            fr = sys._getframe(1)
            if fr.f_code.co_filename in MODULE_FILES:      # calls made by the figure code itself, not by matplotlib
                REC.calls.append(Call(seq=len(REC.calls), ax=self, method=name, args=args, kw=dict(kw),
                                      func=fr.f_code.co_name, lineno=fr.f_lineno, ctx=dict(fr.f_locals),
                                      stack=_stack_funcs(fr), token=REC.token))
        return orig(self, *args, **kw)
    setattr(Axes, name, wrapper)


def _wrap_inset_axes():
    orig = Axes.inset_axes

    @functools.wraps(orig)
    def wrapper(self, *args, **kw):
        child = orig(self, *args, **kw)
        REC.parents[id(child)] = self
        REC.label_axes[id(child)] = child     # keep the inset alive for panel mapping
        return child
    Axes.inset_axes = wrapper


def _wrap_panel_label(mod):
    orig = mod.panel_label

    @functools.wraps(orig)
    def wrapper(ax, s, *a, **k):
        if REC.active:
            REC.labels[id(ax)] = s
            REC.label_axes[id(ax)] = ax
            REC.last_label = s
        REC.suppress += 1             # the letter itself is not data
        try:
            return orig(ax, s, *a, **k)
        finally:
            REC.suppress -= 1
    mod.panel_label = wrapper


def _wrap_panel_function(mod, name):
    """Tag everything done inside one invocation of a panel_* function (reads, marks) with one token."""
    orig = getattr(mod, name)

    @functools.wraps(orig)
    def wrapper(*a, **k):
        if REC.token is not None:
            return orig(*a, **k)
        REC.ntok += 1
        REC.token = (name, REC.ntok)
        try:
            return orig(*a, **k)
        finally:
            REC.token = None
    setattr(mod, name, wrapper)


def _record_read(path):
    if not REC.active:
        return
    try:
        p = Path(path).resolve()
    except TypeError:
        return
    REC.reads.append(dict(path=p, token=REC.token, last_label=REC.last_label, stack=_stack_funcs(sys._getframe(2))))


def _wrap_readers():
    def make(orig):
        @functools.wraps(orig)
        def wrapper(path, *a, **k):
            _record_read(path)
            return orig(path, *a, **k)
        return wrapper
    pd.read_csv = make(READ_CSV)
    pd.read_parquet = make(READ_PARQUET)

    @functools.wraps(READ_TEXT)
    def read_text(self, *a, **k):
        _record_read(self)
        return READ_TEXT(self, *a, **k)
    pathlib.Path.read_text = read_text


def _save(fig, name):
    """Replacement for the modules' save(): the PNG as save() would write it, kept in memory; axes kept."""
    buf = io.BytesIO()
    if hasattr(F, 'fit_width'):          # save() narrows the canvas to the 180 mm journal width first (13_figures.py)
        F.fit_width(fig)
    fig.savefig(buf, format='png', bbox_inches='tight')
    REC.saved.append((name, list(fig.axes), buf.getvalue()))
    plt.close(fig)


_INSTALLED = []


def install():
    if _INSTALLED:
        return
    for n in PLOT_METHODS:
        _wrap_axes_method(n)
    _wrap_inset_axes()
    for mod in (M, F):
        _wrap_panel_label(mod)
        for n in list(vars(mod)):
            if n.startswith('panel_') and n != 'panel_label' and callable(getattr(mod, n)):
                _wrap_panel_function(mod, n)
        mod.save = _save
    _wrap_readers()
    _INSTALLED.append(True)


class FigureRun:
    def __init__(self, key, stem, number, calls, labels, parents, reads, saved):
        self.key, self.stem, self.number = key, stem, number
        self.calls, self.labels, self.parents, self.reads, self.saved = calls, labels, parents, reads, saved
        self.amap = {}
        for _, axes, _ in saved:
            cur = ''
            for ax in axes:
                if id(ax) in labels:
                    cur = labels[id(ax)]
                self.amap[id(ax)] = cur

    def panel_of(self, ax):
        while id(ax) in self.parents:
            ax = self.parents[id(ax)]
        return self.amap.get(id(ax), '')


def run_figure(key, stem, number, fn):
    REC.reset(key)
    M.RNG = np.random.default_rng()
    M.RNG.bit_generator.state = copy.deepcopy(RNG_STATE0)   # as when the figure is drawn by its own invocation
    REC.active = True
    try:
        fn()
    finally:
        REC.active = False
    return FigureRun(key, stem, number, REC.calls, dict(REC.labels), dict(REC.parents), list(REC.reads),
                     list(REC.saved))


# =============================================================================================== small helpers
def num(v):
    if v is None:
        return np.array([], float)
    if isinstance(v, (pd.Series, pd.Index)):
        return v.to_numpy(dtype=float)
    if isinstance(v, range):
        v = list(v)
    return np.atleast_1d(np.asarray(v, dtype=float))


def sidx(v):
    return list(v.index) if isinstance(v, pd.Series) else None


def err_bounds(center, err):
    """Interval bounds from matplotlib xerr/yerr (scalar, symmetric array, or [lower, upper])."""
    c = num(center)
    if err is None:
        return np.full(len(c), np.nan), np.full(len(c), np.nan)
    e = np.asarray([num(x) for x in err], float) if isinstance(err, (list, tuple)) and len(err) == 2 \
        and not np.isscalar(err[0]) else num(err)
    if e.ndim == 1:
        return c - e, c + e
    return c - e[0], c + e[1]


def plain(s):
    """Matplotlib mathtext / LaTeX label -> plain text."""
    if s is None:
        return ''
    s = str(s)
    s = re.sub(r'-\n', '-', s).replace('\n', ' ')
    for a, b in [(r'\geq', '≥'), (r'\leq', '≤'), (r'\times', '×'), (r'\dagger', '†'), (r'\Sigma', 'Σ'),
                 (r'\sum', 'Σ'), (r'\lambda', 'λ'), (r'\mu', 'µ'), (r'\pm', '±'), (r'\leq', '≤'), (r'\,', ' '),
                 (r'\ ', ' '), (r'\%', '%'), (r'\textbar', '|'), ("\\'", "'"), ('~', ' '), ('---', '—'), ('--', '–'),
                 (r'\u2013', '–'), (r'\u2264', '≤')]:
        s = s.replace(a, b)
    s = s.replace('$', '')
    s = re.sub(r'_\{50\}', '50', s)
    s = re.sub(r'\^\{?2\}?', '²', s)
    s = s.replace('^†', '†')
    s = re.sub(r'_\{([^{}]*)\}', r'_\1 ', s)
    s = re.sub(r'\\(?:textbf|textit|emph|texttt|mathrm|text)\{([^{}]*)\}', r'\1', s)
    s = re.sub(r'\\[a-zA-Z]+\s?', '', s)
    s = s.replace('{', '').replace('}', '')
    return re.sub(r'\s+', ' ', s).strip()


def ticks(ax, axis):
    """{position: plain label} of the axis' tick labels as drawn."""
    pos = ax.get_yticks() if axis == 'y' else ax.get_xticks()
    lab = ax.get_yticklabels() if axis == 'y' else ax.get_xticklabels()
    return {float(p): plain(t.get_text()) for p, t in zip(pos, lab)}


def tick_at(ax, axis, v):
    t = ticks(ax, axis)
    if not t:
        return ''
    p = min(t, key=lambda q: abs(q - float(v)))
    return t[p]


def color(c):
    try:
        return to_hex(c)
    except (ValueError, TypeError):
        return str(c)


def legend_entries(calls):
    """(marker, colour, face colour, label) of each legend handle drawn by the panel code."""
    out = []
    for c in calls:
        if c.method != 'legend':
            continue
        for h in c.kw.get('handles', []) or []:
            if isinstance(h, Line2D):
                out.append(dict(marker=str(h.get_marker()), color=color(h.get_color()),
                                mec=color(h.get_markeredgecolor()), mfc=color(h.get_markerfacecolor()),
                                label=plain(h.get_label())))
            else:
                out.append(dict(marker='patch', color=color(h.get_facecolor()), mec=color(h.get_facecolor()),
                                mfc=color(h.get_facecolor()), label=plain(h.get_label())))
    for c in calls:   # empty proxy scatters used for legends
        if c.method == 'scatter' and 'label' in c.kw and len(num(c.args[0])) == 0:
            out.append(dict(marker=str(c.kw.get('marker', 'o')), color=color(c.kw.get('color')),
                            mec=color(c.kw.get('color')), mfc=color(c.kw.get('color')), label=plain(c.kw['label'])))
    return out


def legend_label(entries, marker=None, col=None, mfc=None):
    for e in entries:
        if (marker is None or e['marker'] == marker) and (col is None or color(col) in (e['color'], e['mec'])) \
                and (mfc is None or e['mfc'] == color(mfc)):
            return e['label']
    return ''


def fmt_marker(c):
    fmt = c.kw.get('fmt', c.args[2] if len(c.args) > 2 and isinstance(c.args[2], str) else '')
    return re.sub(r'[-:]', '', fmt or '') or 'o'


class Panel:
    """Recorded marks of one panel."""

    def __init__(self, run, letter, calls):
        self.run, self.letter, self.calls = run, letter, calls
        dc = [c for c in calls if c.method in DATA_METHODS | {'text', 'annotate'}]
        funcs = [c.pfunc for c in dc] or [c.pfunc for c in calls] or ['']
        self.func = max(set(funcs), key=funcs.count)

    def get(self, method=None, func=None, **kw):
        out = []
        for c in self.calls:
            if method and c.method not in (method if isinstance(method, (list, tuple, set)) else [method]):
                continue
            if func and c.pfunc != func and c.func != func:
                continue
            if any(c.kw.get(k) != v for k, v in kw.items()):
                continue
            if c.method in ('scatter', 'plot', 'errorbar') and c.args and len(num(c.args[0])) == 0:
                continue     # empty legend proxies
            out.append(c)
        return out

    @property
    def legend(self):
        return legend_entries(self.calls)

    @property
    def axes(self):
        seen, out = set(), []
        for c in self.calls:
            if id(c.ax) not in seen:
                seen.add(id(c.ax)); out.append(c.ax)
        return out


# =============================================================================================== validation helpers
_UPCACHE = {}


def upstream(path):
    p = PROJECT / path if not Path(path).is_absolute() else Path(path)
    if p not in _UPCACHE:
        if p.suffix == '.json':
            _UPCACHE[p] = json.loads(READ_TEXT(p))
        elif p.suffix == '.parquet':
            _UPCACHE[p] = READ_PARQUET(p)
        else:
            _UPCACHE[p] = READ_CSV(p, low_memory=False)
    return _UPCACHE[p]


def _cmp(a, b):
    try:
        a, b = float(a), float(b)
    except (TypeError, ValueError):
        return (str(a) == str(b)), np.nan
    if np.isnan(a) and np.isnan(b):
        return True, 0.0
    return a == b, abs(a - b)


def check(df, col, path, keys, up_col, up_filter=None, transform=None, mask=None, note=''):
    """Exact comparison of df[col] with the upstream value found by keys [(df column, upstream column)]."""
    U = upstream(path)
    if up_filter is not None:
        U = up_filter(U)
    d = df if mask is None else df[mask]
    d = d[d[col].notna()]
    n = ok = 0
    worst = 0.0
    missing = 0
    for _, r in d.iterrows():
        q = U
        for a, b in keys:
            q = q[q[b] == r[a]]
        if len(q) != 1:
            missing += 1
            continue
        v = q[up_col].iloc[0]
        if transform is not None:
            v = transform(v)
        e, diff = _cmp(r[col], v)
        n += 1; ok += e
        if np.isfinite(diff):
            worst = max(worst, diff)
    res = 'PASS' if n and ok == n and not missing else 'FAIL'
    return dict(check='exact lookup', exported_column=col, upstream_file=str(path), upstream_column=up_col,
                key=', '.join(f'{a}={b}' for a, b in keys), n_checked=n, n_exact_match=ok, n_unmatched_key=missing,
                max_abs_diff=worst, result=res, note=note)


def check_values(name, exported, expected, path, up_col, note=''):
    """Exact comparison of two equally long value sequences (exported vs computed from the upstream file)."""
    a, b = list(exported), list(expected)
    n = min(len(a), len(b))
    res = [_cmp(x, y) for x, y in zip(a, b)]
    ok = sum(e for e, _ in res)
    worst = max([d for _, d in res if np.isfinite(d)] or [0.0])
    return dict(check='exact recomputation' if 'recomput' in note else 'exact comparison', exported_column=name,
                upstream_file=str(path), upstream_column=up_col, key='', n_checked=n, n_exact_match=ok,
                n_unmatched_key=abs(len(a) - len(b)), max_abs_diff=worst,
                result='PASS' if n and ok == n and len(a) == len(b) else 'FAIL', note=note)


def ci_pair(s):
    return [float(x) for x in str(s).strip('[]() ').split(',')]


# =============================================================================================== converters
CONVERTERS = {}


def conv(*keys):
    def deco(fn):
        for k in keys:
            CONVERTERS[k] = fn
        return fn
    return deco


def T(rows, cols):
    """DataFrame with the described columns first, in order."""
    df = pd.DataFrame(rows)
    order = [c for c in cols if c in df.columns] + [c for c in df.columns if c not in cols]
    return df[order] if len(df) else pd.DataFrame(columns=list(cols))


# ---------------------------------------------------------------------------------------------- Figure 1
@conv('panel_schematic')
def c_schematic(P):
    rows = []
    for c in P.get('imshow'):
        m = np.asarray(c.args[0])
        for i in range(m.shape[0]):
            for j in range(m.shape[1]):
                rows.append(dict(matrix=plain(c.ctx.get('t')), drug_row=i + 1, cell_line_column=j + 1,
                                 value=float(m[i, j])))
    cols = OrderedDict(matrix='which of the three schematic matrices (profiles z = loading × g + selective residual)',
                       drug_row='row (illustrative drug) of the 6 × 12 matrix',
                       cell_line_column='column (illustrative cell line), cell lines sorted by g',
                       value='matrix entry (colour scale RdBu_r, -2.5 to 2.5)')
    note = ('Schematic: synthetic values drawn in the figure code (numpy default_rng(1)); no measured data. '
            'Exported so that the panel can be redrawn.')
    return T(rows, cols), cols, [], note


@conv('panel_axis_share')
def c_axis_share(P):
    c = P.get('boxplot')[0]
    lab = ticks(c.ax, 'x')
    stats = c.ctx['stats']
    inv = {v: k for k, v in F.LABEL.items()}
    rows = []
    for k, d in enumerate(c.args[0]):
        screen = lab.get(float(k + 1), str(k + 1))
        cand = stats[stats.screen == inv.get(screen, screen)]
        vc = cand.fitness_r2.value_counts()
        look = dict(zip(cand.fitness_r2, cand.compound))
        for v in d:
            rows.append(dict(screen=screen, compound=look[v] if vc.get(v, 0) == 1 else '', fitness_r2=float(v)))
    df = T(rows, [])
    cols = OrderedDict(screen='discovery screen (box)',
                       compound='drug (harmonized compound id), identified from the panel input by exact value',
                       fitness_r2="share of the drug's within-screen profile z explained by g (R², fitness_r2); "
                                  "active drugs (≥10% of cell lines with AUC < 0.8). Box: quartiles; whiskers 1.5 IQR; "
                                  "number: median")
    df = df[list(cols)]
    chk = [check(df, 'fitness_r2', 'results/decomposition/drug_screen_stats.csv',
                  [('screen', 'screen'), ('compound', 'compound')], 'fitness_r2')]
    return df, cols, chk, ''


@conv('panel_ladder')
def c_ladder(P):
    rows = []
    for c in P.get('scatter'):
        d = c.ctx['d']
        idx = sidx(c.args[0])
        for i, v in zip(idx, num(c.args[0])):
            rows.append(dict(relationship=plain(c.ctx['lab']), endpoint=plain(c.ctx['fam']).replace('lnIC50', 'ln IC50'),
                             drug_set=c.ctx['v'], comparison=d.loc[i, 'label'], axis_agreement_r=v,
                             mark='filled circle' if c.ctx['filled'] else 'open circle'))
    for c in P.get('plot'):
        rows.append(dict(relationship=plain(c.ctx['lab']), endpoint=plain(c.ctx['fam']).replace('lnIC50', 'ln IC50'),
                         drug_set='all', comparison='median over the comparisons above (all drugs)',
                         axis_agreement_r=num(c.args[0])[0], mark='bar'))
    cols = OrderedDict(relationship='relationship between the two measurements (row of the panel)',
                       endpoint='endpoint family (AUC, orange; ln IC50, aqua)',
                       drug_set='axes built from all drugs (open) or active drugs only (filled)',
                       comparison='the two datasets or halves compared (dataset name: split-half within that dataset)',
                       axis_agreement_r='Pearson r between the two general-sensitivity axes over shared cell lines',
                       mark='mark in the panel')
    df = T(rows, cols)
    pts = df[df.mark != 'bar']
    ds = upstream('results/revision_axis/dreval_datasets.csv')
    wl = upstream('results/revision_axis/within_lab_controls.csv')
    ap = upstream('results/revision_axis/all_axis_pairs.csv')
    pool = set(ds.split_half) | set(ds.split_half_active) | set(wl.r) | set(ap.r)
    chk = [dict(check='exact membership', exported_column='axis_agreement_r (points)',
                upstream_file='results/revision_axis/{dreval_datasets,within_lab_controls,all_axis_pairs}.csv',
                upstream_column='split_half, split_half_active, r', key='', n_checked=len(pts),
                n_exact_match=int(pts.axis_agreement_r.isin(pool).sum()), n_unmatched_key=0, max_abs_diff=0.0,
                result='PASS' if len(pts) and pts.axis_agreement_r.isin(pool).all() else 'FAIL',
                note='every plotted r is a value of the upstream files')]
    oth = df[(df.mark != 'bar') & df.relationship.str.startswith('other institution')].copy()
    oth['a'] = oth.comparison.str.split('-').str[0]
    oth['b'] = oth.comparison.str.split('-').str[1]
    oth['ep'] = oth.endpoint.map({'AUC': 'CC_AUC', 'ln IC50': 'CC_lnIC50'})
    chk.append(check(oth, 'axis_agreement_r', 'results/revision_axis/all_axis_pairs.csv',
                     [('a', 'a'), ('b', 'b'), ('ep', 'endpoint'), ('drug_set', 'variant')], 'r',
                     up_filter=lambda U: U[U.source == '46 DrEval'], note='other-institution rows, DrEval CurveCurator'))
    return df, cols, chk, ''


@conv('panel_shared_fraction')
def c_shared_fraction(P):
    rows = []
    for c in P.get('errorbar'):
        r = c.ctx['r']
        rows.append(dict(axis_set=c.ctx['lab'], axis_set_definition=c.ctx['key'], screen=F.LABEL.get(c.ctx['s'], c.ctx['s']),
                         screen_code=c.ctx['s'], lambda2=num(c.args[1])[0], ci_lo=r.ci_lo, ci_hi=r.ci_hi,
                         triads=r.triads))
    cols = OrderedDict(axis_set='axes built from all, active or inactive drugs (legend)',
                       axis_set_definition='axis_set as named in the upstream file',
                       screen='discovery screen', screen_code='screen code in the upstream file',
                       lambda2='share of the screen axis loading on the factor common to all screens (λ², median over '
                               'triads of screens)',
                       ci_lo='95% cell-bootstrap interval, lower', ci_hi='95% cell-bootstrap interval, upper',
                       triads='number of screen triads')
    df = T(rows, cols)
    chk = [check(df, c, 'results/revision_axis/triad_shared_fraction.csv',
                 [('axis_set_definition', 'axis_set'), ('screen_code', 'screen')], c) for c in ('lambda2', 'ci_lo')]
    return df, cols, chk, ''


@conv('panel_variance')
def c_variance(P):
    rows = []
    for c in P.get('barh'):
        rows.append(dict(data=plain(c.ctx['lab']), model=c.ctx['m'], component=c.ctx['clab'], component_code=c.ctx['c'],
                         fraction=num(c.args[1])[0], bar_left=float(c.kw.get('left', 0.0))))
    cols = OrderedDict(data='data set and processing (row of the panel)', model='REML model as named in the upstream file',
                       component='variance component (legend): shared by all screens, institution, screen, '
                                 'screen/run/site (residual)',
                       component_code='component as named in the upstream file',
                       fraction='share of the cell-level variance of the axis (bar length)',
                       bar_left='cumulative share to the left of the bar (stacking)')
    df = T(rows, cols)
    chk = [check(df, 'fraction', 'results/revision_axis/variance_components.csv',
                 [('model', 'model'), ('component_code', 'component')], 'fraction')]
    return df, cols, chk, ''


# ---------------------------------------------------------------------------------------------- Figure 2
@conv('panel_leaderboard')
def c_leaderboard(P):
    rows = []
    for c in P.get('errorbar'):
        r = c.ctx['r']
        rows.append(dict(model=r.model, model_label=plain(M.NAMES.get(r.model, r.model)), model_class=r.klass,
                         drug_agnostic=fmt_marker(c) == '^', drugs=r.drugs, median_per_drug_r=num(c.args[0])[0],
                         per_drug_ci_lo=r.raw_lo, per_drug_ci_hi=r.raw_hi, median_selective_r=num(c.args[1])[0],
                         selective_ci_lo=r.selective_lo, selective_ci_hi=r.selective_hi, element='model'))
    for c in P.get('axhspan'):
        rows.append(dict(model='drug-agnostic floor', model_label='drug-agnostic floor (shaded)',
                         median_selective_r=num(c.args[1])[0], element='reference: best median selective correlation '
                                                                        'of a drug-agnostic predictor'))
    for c in P.get('axvline'):
        rows.append(dict(model='fitness_measured', model_label='measured axis g (dashed line)',
                         median_per_drug_r=num(c.args[0])[0], element='reference: measured axis g, not cell-blind'))
    cols = OrderedDict(model='model (DrEval name, or this study)', model_label='name in the figure',
                       model_class='symbol colour: naive, classical ML, neural/factorization, this study',
                       drug_agnostic='triangle: drug-agnostic predictor',
                       drugs='drugs scored', median_per_drug_r='median over drugs of the per-drug correlation (x)',
                       per_drug_ci_lo='95% drug-bootstrap interval, lower', per_drug_ci_hi='95% drug-bootstrap interval, upper',
                       median_selective_r='median over drugs of the selective correlation (y)',
                       selective_ci_lo='95% drug-bootstrap interval, lower',
                       selective_ci_hi='95% drug-bootstrap interval, upper', element='model or reference line')
    df = T(rows, cols)
    m = df.element == 'model'
    chk = [check(df, c, 'results/dreval/model_summary.csv', [('model', 'model')], u, mask=m)
           for c, u in [('median_per_drug_r', 'raw'), ('median_selective_r', 'selective'), ('selective_ci_lo', 'selective_lo')]]
    chk.append(check(df, 'median_per_drug_r', 'results/dreval/model_summary.csv', [('model', 'model')], 'raw',
                     mask=df.model == 'fitness_measured', note='measured axis line'))
    return df, cols, chk, ''


def _boot_note():
    return ('Intervals are drug-bootstrap intervals computed in the figure script (64_figures_revision.py boot_means, '
            '1000 resamples, module RNG seeded 20261015 at the start of the figure); they reproduce the published '
            'figure.')


def _per_drug_means(pair, cols, models):
    P = READ_CSV(PROJECT / 'results/revision_crosslab/per_drug_scores.csv.gz', low_memory=False,
                 usecols=['pair', 'model', 'drug', 'class'] + cols)
    P = P[P.pair == pair]
    out = {}
    for m in models:
        X = P[P.model == m][cols].to_numpy(float)
        X = X[np.isfinite(X).all(1)]
        out[m] = [X[:, j].mean() for j in range(X.shape[1])]     # as in boot_means()
    return out


@conv('panel_decomposition')
def c_decomposition(P):
    b1 = [c for c in P.get('barh') if 'left' not in c.kw][0]
    b2 = [c for c in P.get('barh') if 'left' in c.kw][0]
    e = P.get('errorbar')[0]
    B = b1.ctx['B']
    lab = ticks(b1.ax, 'y')
    rows = []
    for y, i, axis, sel_bar, raw in zip(num(b1.args[0]), sidx(b1.args[1]), num(b1.args[1]), num(b2.args[1]),
                                        num(e.args[0])):
        r = B.loc[i]
        rows.append(dict(model=r.model, model_label=lab.get(float(y), ''), drugs=r.drugs, mean_per_drug_r=raw,
                         mean_per_drug_r_ci_lo=r.raw_w_lo, mean_per_drug_r_ci_hi=r.raw_w_hi, axis_term=axis,
                         axis_term_ci_lo=r.med_w_lo, axis_term_ci_hi=r.med_w_hi, selective_term=r.selterm_w,
                         selective_term_ci_lo=r.selterm_w_lo, selective_term_ci_hi=r.selterm_w_hi,
                         selective_bar_drawn=sel_bar, pct_through_axis=100 * r.share,
                         pct_through_axis_ci_lo=100 * r.share_lo, pct_through_axis_ci_hi=100 * r.share_hi))
    cols = OrderedDict(model='model', model_label='name in the figure (rows sorted by mean_per_drug_r)',
                       drugs='drugs with finite terms',
                       mean_per_drug_r='mean over drugs of the per-drug correlation in CTRPv2 (error bar centre)',
                       mean_per_drug_r_ci_lo='95% drug-bootstrap interval, lower',
                       mean_per_drug_r_ci_hi='95% drug-bootstrap interval, upper',
                       axis_term='mean axis-mediated term r_pg·r_zg (grey bar)',
                       axis_term_ci_lo='95% drug-bootstrap interval, lower', axis_term_ci_hi='95% drug-bootstrap interval, upper',
                       selective_term='mean selective term (mean_per_drug_r - axis_term)',
                       selective_term_ci_lo='95% drug-bootstrap interval, lower',
                       selective_term_ci_hi='95% drug-bootstrap interval, upper',
                       selective_bar_drawn='orange bar as drawn (selective term clipped at 0)',
                       pct_through_axis='percentage of the mean per-drug correlation carried by the axis (number in '
                                        'the panel)',
                       pct_through_axis_ci_lo='95% drug-bootstrap interval, lower',
                       pct_through_axis_ci_hi='95% drug-bootstrap interval, upper')
    df = T(rows, cols)
    mu = _per_drug_means(M.PAIR_WITHIN, ['raw_w', 'med_w', 'selterm_w'], list(df.model))
    chk = [check_values('mean_per_drug_r', df.mean_per_drug_r, [mu[m][0] for m in df.model],
                        'results/revision_crosslab/per_drug_scores.csv.gz', 'raw_w',
                        note=f'recomputation: mean over drugs (pair {M.PAIR_WITHIN}, rows with finite terms)'),
           check_values('axis_term', df.axis_term, [mu[m][1] for m in df.model],
                        'results/revision_crosslab/per_drug_scores.csv.gz', 'med_w',
                        note='recomputation: mean over drugs')]
    return df, cols, chk, _boot_note()


@conv('panel_paired')
def c_paired(P):
    rows = []
    for c in P.get('scatter'):
        B, a, b = c.ctx['B'], c.ctx['a'], c.ctx['b']
        for i, x, y in zip(sidx(c.args[0]), num(c.args[0]), num(c.args[1])):
            r = B.loc[i]
            rows.append(dict(model=r.model, term=c.ctx['lab'], pair=c.ctx['pair'], drugs=r.drugs,
                             mean_term_scored_in_training=x, training_ci_lo=r[a + '_lo'], training_ci_hi=r[a + '_hi'],
                             mean_term_scored_in_target=y, target_ci_lo=r[b + '_lo'], target_ci_hi=r[b + '_hi']))
    cols = OrderedDict(model='published DrEval model (classical and neural/factorization, SVR excluded)',
                       term='axis term (grey) or selective term (orange)',
                       pair='training dataset -> target dataset (CTRPv2 CurveCurator ln IC50 -> GDSC1 CurveCurator ln IC50)',
                       drugs='drugs with finite terms',
                       mean_term_scored_in_training='mean over drugs of the term, predictions scored in CTRPv2 (x)',
                       training_ci_lo='95% drug-bootstrap interval, lower', training_ci_hi='95% drug-bootstrap interval, upper',
                       mean_term_scored_in_target='mean over drugs of the term, same predictions scored in GDSC1 (y)',
                       target_ci_lo='95% drug-bootstrap interval, lower', target_ci_hi='95% drug-bootstrap interval, upper')
    df = T(rows, cols)
    pair = df.pair.iloc[0]
    mu = _per_drug_means(pair, ['med_w', 'med_c', 'selterm_w', 'selterm_c'], list(df.model.unique()))
    ax = df[df.term == 'axis term']
    chk = [check_values('mean_term_scored_in_target (axis term)', ax.mean_term_scored_in_target,
                        [mu[m][1] for m in ax.model], 'results/revision_crosslab/per_drug_scores.csv.gz', 'med_c',
                        note=f'recomputation: mean over drugs (pair {pair}, rows with finite terms)')]
    return df, cols, chk, _boot_note() + ' Lines: identity (fully retained) and half retained.'


def _target_checks(df, col, up_col, transform=None):
    out = []
    for f in ['results/revision_crosslab/summary_by_target.csv',
              'results/revision_crosslab/harmonized_range/summary_by_target.csv']:
        U = upstream(f)
        U = U[(U.subset == 'all') & ((U.cellset == 'own') if 'cellset' in U and 'harmonized' not in f else True)]
        m = df.pair.isin(set(U.pair))
        if m.any():
            out.append(check(df, col, f, [('pair', 'pair')], up_col, transform=transform, mask=m,
                             up_filter=lambda V, f=f: V[(V.subset == 'all') & ((V.cellset == 'own') if 'harmonized' not in f
                                                                              else True)]))
    return out


@conv('panel_retention')
def c_retention(P):
    leg = P.legend
    rows = []
    for c in P.get('errorbar'):
        r = c.ctx['r']
        rows.append(dict(target=plain(r.label), group=r.group, pair=r.pair, drugs=r.drugs,
                         quantity=legend_label(leg, marker=fmt_marker(c)), ratio=num(c.args[0])[0],
                         ci_lo=c.ctx['lo'], ci_hi=c.ctx['hi']))
    cols = OrderedDict(target='target dataset and processing version (row label)', group='target dataset',
                       pair='training -> target dataset|processing|endpoint, as in the upstream files',
                       drugs='drugs (in parentheses in the row label)',
                       quantity='axis term r_pg·r_zg retained (circle), alignment r_pg (square), target loading r_zg '
                                '(diamond)',
                       ratio='ratio, other dataset / CTRPv2 (pooled over 21 published models)',
                       ci_lo='95% two-way (drug × cell line) bootstrap interval, lower', ci_hi='95% interval, upper')
    df = T(rows, cols)
    chk = []
    for q, up in [('axis term', 'mediated_retained'), ('alignment', 'alignment_ratio_r_pg'),
                  ('loading', 'target_loading_ratio_r_zg')]:
        sub = df[df.quantity.str.startswith(q)].reset_index(drop=True)
        chk += _target_checks(sub, 'ratio', up)
    return df, cols, chk, ''


@conv('panel_change')
def c_change(P):
    rows = OrderedDict()
    for c in P.get('errorbar'):
        r = c.ctx['r']
        d = rows.setdefault(r.pair, dict(target=plain(r.label), group=r.group, pair=r.pair, drugs=r.drugs))
        if fmt_marker(c) == 'd':
            d.update(change_per_drug_raw=num(c.args[0])[0], raw_ci95_lo=r.draw_lo, raw_ci95_hi=r.draw_hi)
        else:
            d.update(change_selective_both=num(c.args[0])[0], selboth_ci90_lo=r.dsel_lo, selboth_ci90_hi=r.dsel_hi,
                     shown_to_retain_80pct=r.kept, open_symbol=color(c.kw.get('mfc')) == '#ffffff',
                     leakage_floor=r.floor, leakage_floor_highlighted=bool(r.floor >= 0.05))
    cols = OrderedDict(target='target dataset and processing version (row)', group='target dataset',
                       pair='training -> target, as in the upstream files', drugs='drugs',
                       change_per_drug_raw='change in the mean per-drug correlation, other dataset - CTRPv2 (grey)',
                       raw_ci95_lo='95% interval, lower', raw_ci95_hi='95% interval, upper',
                       change_selective_both='change in the symmetric selective-both score, other dataset - CTRPv2 '
                                             '(orange)',
                       selboth_ci90_lo='90% interval, lower', selboth_ci90_hi='90% interval, upper',
                       shown_to_retain_80pct='retained-fraction test: one-sided 95% bound ≥ 0.8 (False: open symbol; '
                                             'blank: not tested for this row)',
                       open_symbol='open symbol in the panel',
                       leakage_floor='best median selective-both of a drug-agnostic predictor in that target (number '
                                     'at right; printed clipped at 0)',
                       leakage_floor_highlighted='printed in orange (floor ≥ 0.05)')
    df = T(list(rows.values()), cols)
    chk = _target_checks(df, 'change_per_drug_raw', 'raw_loss', transform=lambda v: -v)
    Y = 'results/revision_round2/symmetric_selboth/selboth_by_target.csv'
    chk.append(check(df, 'change_selective_both', Y, [('pair', 'pair')], 'change_sym',
                     up_filter=lambda U: U[(U.benchmark == 'leaderboard') & (U.group == 'main')]))
    return df, cols, chk, ''


# ---------------------------------------------------------------------------------------------- Figure 3
@conv('panel_improve_share')
def c_improve_share(P):
    rows = []
    for c in P.get('scatter'):
        x = c.ctx['x']
        for i, v in zip(sidx(c.args[1]), num(c.args[1])):
            rows.append(dict(dataset=c.ctx['d'], series='model', model=x.loc[i, 'model'], axis_share=v))
    for c in P.get('errorbar'):
        b = c.ctx['b']
        rows.append(dict(dataset=c.ctx['d'], series='pooled over the 7 models', model='pooled',
                         axis_share=num(c.args[1])[0], ci_lo=b.lo, ci_hi=b.hi))
    for c in P.get('plot'):
        rows.append(dict(dataset=c.ctx['d'], series='perfect predictor (mean r_zg²)', model='perfect predictor',
                         axis_share=num(c.args[1])[0]))
    cols = OrderedDict(dataset='IMPROVE dataset', series='light dots: models; dark: pooled; orange: perfect predictor',
                       model='IMPROVE model', axis_share='axis share of per-drug accuracy (Σ axis term / Σ raw correlation '
                                                         'over drugs)',
                       ci_lo='two-way (drug × cell) bootstrap 95% interval, lower', ci_hi='95% interval, upper')
    df = T(rows, cols)
    W = 'results/revision_improve/no_split_replicates/within_summary.csv'
    chk = [check(df, 'axis_share', W, [('dataset', 'dataset'), ('model', 'model')], 'axis_share', mask=df.series == 'model'),
           check(df, 'axis_share', 'results/revision_improve/no_split_replicates/within_pooled_bootstrap.csv',
                 [('dataset', 'dataset')], 'estimate', mask=df.model == 'pooled',
                 up_filter=lambda U: U[U.stat == 'fitness_share_source'])]
    return df, cols, chk, ''


@conv('panel_improve_floor')
def c_improve_floor(P):
    rows = []
    for c in P.get('scatter'):
        if c.kw.get('marker') == '^':
            rows.append(dict(dataset=c.ctx['d'], series='drug-agnostic floor (cell-line mean)', model='floor_cell_mean',
                             median_per_drug_raw_r=num(c.args[1])[0]))
        else:
            x = c.ctx['x']
            for i, v in zip(sidx(c.args[1]), num(c.args[1])):
                rows.append(dict(dataset=c.ctx['d'], series='model', model=x.loc[i, 'model'], median_per_drug_raw_r=v))
    cols = OrderedDict(dataset='IMPROVE dataset', series='model (dots) or drug-agnostic floor (triangle and bar)',
                       model='IMPROVE model', median_per_drug_raw_r='median over drugs of the per-drug raw correlation')
    df = T(rows, cols)
    chk = [check(df, 'median_per_drug_raw_r', 'results/revision_improve/no_split_replicates/within_summary.csv',
                 [('dataset', 'dataset'), ('model', 'model')], 'median_raw')]
    return df, cols, chk, ''


@conv('panel_improve_transfer')
def c_improve_transfer(P):
    rows = OrderedDict()
    for c in P.get('scatter'):
        S = c.ctx['S']
        src = color(c.kw.get('facecolor')) == '#ffffff'
        for i, x, y in zip(sidx(c.args[0]), num(c.args[0]), num(c.args[1])):
            r = S.loc[i]
            d = rows.setdefault((r.source, r.target), dict(source=r.source, target=r.target, drugs=r.drugs,
                                                           measurement_selective_both_in_target=x))
            d['models_selective_both_in_source' if src else 'models_selective_both_in_target'] = y
    for c in P.get('plot', alpha=0.6):
        r, q = c.ctx['r'], c.ctx['q']
        d = rows[(r.source, r.target)]
        d.update(target_ci_lo=num(c.args[1])[0], target_ci_hi=num(c.args[1])[1],
                 ratio_models_to_measurement=q.ratio_to_measurement)
    cols = OrderedDict(source='training (source) dataset', target='dataset the identical predictions are scored in',
                       drugs='drugs (symbol size)',
                       measurement_selective_both_in_target="selective-both of the source's own measurement scored in the "
                                                            "target (x)",
                       models_selective_both_in_source='selective-both of the predictions in the source (open), mean over '
                                                       'models',
                       models_selective_both_in_target='selective-both of the identical predictions in the target '
                                                       '(filled), mean over models',
                       target_ci_lo='two-way bootstrap 95% interval of the pooled models in the target, lower',
                       target_ci_hi='95% interval, upper',
                       ratio_models_to_measurement='models / measurement (ratio summarized in the panel text)')
    df = T(list(rows.values()), cols)
    S = 'results/revision_improve/no_split_replicates/selective_both_change.csv'
    f = lambda U: U[U.axis_panel == 'own_panel']
    chk = [check(df, 'models_selective_both_in_target', S, [('source', 'source'), ('target', 'target')],
                 'models_cross_mean_same_drugs', up_filter=f),
           check(df, 'measurement_selective_both_in_target', S, [('source', 'source'), ('target', 'target')],
                 'measurement_cross_mean', up_filter=f)]
    return df, cols, chk, 'Dashed line: identity.'


ENDPOINT_NAME = {'lnic50': 'ln IC50 (CurveCurator)', 'auc': 'AUC (CurveCurator)'}


@conv('panel_axis_weight')
def c_axis_weight(P):
    rows = []
    for c in P.get('scatter'):
        d = c.ctx['d']
        imp = c.kw.get('marker') == '^'
        rows.append(dict(dataset=d, dataset_label=tick_at(c.ax, 'y', c.ctx['k']),
                         endpoint="IMPROVE fixed-range AUC" if imp else ENDPOINT_NAME.get(c.ctx['ep'], c.ctx['ep']),
                         endpoint_code=M.RT_IM.get(d) if imp else c.ctx['ep'], source='IMPROVE' if imp else 'DrEval',
                         perfect_predictor_axis_share=num(c.args[0])[0]))
    cols = OrderedDict(dataset='dataset', dataset_label='row label', endpoint='endpoint (filled: ln IC50; open: AUC; '
                                                                             'triangle: IMPROVE)',
                       endpoint_code='endpoint (DrEval) or IMPROVE dataset name in the upstream file',
                       source='benchmark', perfect_predictor_axis_share="perfect predictor's axis share (mean over drugs "
                                                                       "of r_zg²)")
    df = T(rows, cols)
    chk = [check(df, 'perfect_predictor_axis_share', 'results/revision_retrain/dataset_fitness.csv',
                 [('dataset', 'dataset'), ('endpoint_code', 'endpoint')], 'mean_fitness_r2', mask=df.source == 'DrEval'),
           check(df, 'perfect_predictor_axis_share', 'results/revision_improve/no_split_replicates/within_summary.csv',
                 [('endpoint_code', 'dataset')], 'oracle_share', mask=df.source == 'IMPROVE',
                 up_filter=lambda U: U.groupby('dataset', as_index=False).oracle_share.first())]
    return df, cols, chk, ''


@conv('panel_capture')
def c_capture(P):
    rows, tick = [], {}
    for c in P.get('scatter', marker='_'):
        d, ok = c.ctx['d'], c.ctx['ok']
        for m, y in zip(d.model.values[ok], num(c.args[1])):
            tick[(c.ctx['lab'], m)] = y
    for c in P.get('scatter'):
        mk = c.kw.get('marker', 'o')
        if mk == '_':
            continue
        if mk == 'x':
            Td = c.ctx['T']
            for i, x, y in zip(sidx(c.args[0]), num(c.args[0]), num(c.args[1])):
                r = Td.loc[i]
                rows.append(dict(set='tissue-by-drug mean', benchmark=r.benchmark, setting=r.setting, model=r.model,
                                 axis_capture=x, selective_capture=y, marker='x (purple)'))
            continue
        d, lab = c.ctx['d'], c.ctx['lab']
        for i, x, y in zip(sidx(c.args[0]), num(c.args[0]), num(c.args[1])):
            r = d.loc[i]
            rows.append(dict(set=plain(lab), benchmark=r.benchmark, setting=r.setting, model=r.model, axis_capture=x,
                             selective_capture=y, selective_capture_given_tissue=tick.get((lab, r.model), np.nan),
                             marker=f"{mk} ({'open' if color(c.kw.get('facecolor')) == '#ffffff' else 'filled'})"))
    for c in P.get('plot'):
        if c.kw.get('marker') == '*':      # the pooled capture of the published models (star and its label)
            r = c.ctx['pool']
            rows.append(dict(set=f'pooled over the {int(r.models)} published models (star)', benchmark=r.benchmark,
                             setting=r.setting, model=r.group, axis_capture=num(c.args[0])[0],
                             selective_capture=num(c.args[1])[0], marker='* (blue)'))
    cols = OrderedDict(set='group of models (legend): released leaderboard models, drug-agnostic, retrained in drevalpy '
                           '(by training dataset and endpoint), tissue-by-drug mean, pooled published models (star)',
                       benchmark='benchmark in the upstream file', setting='training dataset and endpoint',
                       model='model', axis_capture="fraction of a perfect predictor's axis term recovered (x)",
                       selective_capture="fraction of a perfect predictor's selective term recovered (y)",
                       selective_capture_given_tissue='selective capture when tissue is also partialled out (tick)',
                       marker='marker in the panel')
    df = T(rows, cols)
    C = 'results/revision_round2/capture/capture_by_model.csv'
    L = 'results/revision_round2/lineage/lineage_by_model.csv'
    k = [('benchmark', 'benchmark'), ('setting', 'setting'), ('model', 'model')]
    pooled = df.set.str.startswith('pooled')
    m = (df.set != 'tissue-by-drug mean') & ~pooled
    CP = 'results/revision_round2/capture/capture_pooled.csv'
    kp = [('benchmark', 'benchmark'), ('setting', 'setting'), ('model', 'group')]
    chk = [check(df, 'axis_capture', CP, kp, 'axis_capture', mask=pooled),
           check(df, 'selective_capture', CP, kp, 'selective_capture', mask=pooled),
           check(df, 'axis_capture', C, k, 'axis_capture', mask=m),
           check(df, 'selective_capture', C, k, 'selective_capture', mask=m),
           check(df, 'selective_capture_given_tissue', L, k, 'selective_capture_lin', mask=m,
                 up_filter=lambda U: U[U.kind == 'learned']),
           check(df, 'axis_capture', L, k, 'axis_capture', mask=df.set == 'tissue-by-drug mean',
                 up_filter=lambda U: U[U.kind == 'tissue-by-drug baseline'])]
    return df, cols, chk, ('Dashed line: equal capture. Dotted line: axis capture 1 (the perfect predictor\'s axis term). '
                           'The perfect predictor is at (1, 1), outside the plotted range. The star\'s label gives its '
                           'two values as percentages.')


@conv('panel_retrain_terms')
def c_retrain_terms(P):
    leg = P.legend
    rows = []
    for c in P.get('errorbar'):
        r = c.ctx['r']
        rows.append(dict(training=plain(M.RT_LABEL.get(c.ctx['st'], c.ctx['st'])), setting=c.ctx['st'],
                         target=c.ctx['t'], term=legend_label(leg, marker=fmt_marker(c)), column=c.ctx['col'],
                         change=num(c.args[0])[0], ci_lo=c.ctx['lo'], ci_hi=c.ctx['hi'], models=r.models))
    cols = OrderedDict(training='training dataset and endpoint', setting='setting in the upstream file',
                       target='dataset the identical predictions are scored in',
                       term='axis term (grey diamond) or selective term (orange circle)',
                       column='quantity in the upstream file',
                       change='change in the mean per-drug term, other dataset - training dataset, pooled over models',
                       ci_lo='two-way bootstrap 95% interval, lower', ci_hi='95% interval, upper', models='models pooled')
    df = T(rows, cols)
    chk = [check(df, 'change', 'results/revision_retrain/paired_pooled.csv', [('setting', 'setting'), ('target', 'target')],
                 col, mask=df.column == col, up_filter=lambda U: U.drop_duplicates(['setting', 'target']))
           for col in ('delta_mediated', 'delta_selective_term')]
    return df, cols, chk, ''


@conv('panel_retrain_transfer')
def c_retrain_transfer(P):
    rows = OrderedDict()
    for c in P.get(['scatter', 'plot']):
        st, t = c.ctx.get('st'), c.ctx.get('t')
        if st is None or t is None or 'm' not in c.ctx:
            continue
        S = c.ctx['S']
        d = rows.setdefault((st, t), dict(training=plain(M.RT_LABEL.get(st, st)), setting=st, target=t,
                                          drugs=S[(S.setting == st) & (S.target == t)].drugs.iloc[0],
                                          measurement_selective_both_in_target=c.ctx['m']))
        mk = c.kw.get('marker', 'o')
        if c.method == 'scatter' and mk in ('+', 'x'):
            d['tissue_by_drug_marker_x'] = num(c.args[0])[0]
        if c.method == 'scatter' and mk == '+':
            d['tissue_by_drug_mean_in_training'] = num(c.args[1])[0]
        elif c.method == 'scatter' and mk == 'x':
            d['tissue_by_drug_mean_in_target'] = num(c.args[1])[0]
        elif c.method == 'scatter' and color(c.kw.get('facecolor')) == '#ffffff':
            d['models_selective_both_in_training'] = num(c.args[1])[0]
        elif c.method == 'scatter':
            d['models_selective_both_in_target'] = num(c.args[1])[0]
        elif c.method == 'plot' and c.kw.get('alpha') == 0.6 and 'q' in c.ctx and c.kw.get('lw') == 0.6:
            d['target_ci_lo'], d['target_ci_hi'] = num(c.args[1])
    cols = OrderedDict(training='training dataset and endpoint (colour)', setting='setting in the upstream file',
                       target='dataset the identical predictions are scored in', drugs='drugs (symbol size)',
                       measurement_selective_both_in_target="selective-both of the training dataset's own measurement in "
                                                            "the target (x)",
                       models_selective_both_in_training='pooled selective-both of the predictions in the training '
                                                         'dataset (open)',
                       models_selective_both_in_target='pooled selective-both of the identical predictions in the target '
                                                       '(filled)',
                       target_ci_lo='two-way bootstrap 95% interval (target), lower', target_ci_hi='95% interval, upper',
                       tissue_by_drug_mean_in_training="DrEval's tissue-by-drug mean, selective-both in the training "
                                                       "dataset (+)",
                       tissue_by_drug_mean_in_target="DrEval's tissue-by-drug mean, selective-both in the target (x)",
                       tissue_by_drug_marker_x='x position of the + and x markers as drawn (the measurement value '
                                               'shifted left by 0.008 so the markers do not overlap)')
    df = T(list(rows.values()), cols)
    k = [('setting', 'setting'), ('target', 'target')]
    chk = [check(df, 'models_selective_both_in_target', 'results/revision_retrain/paired_pooled.csv', k, 'mean_sel_BOTH_t',
                 up_filter=lambda U: U.drop_duplicates(['setting', 'target'])),
           check(df, 'measurement_selective_both_in_target', 'results/revision_retrain/paired_summary.csv', k,
                 'mean_sel_BOTH_t', up_filter=lambda U: U[(U.panel == 'main') & (U.model == 'ref:source_measurement')])]
    return df, cols, chk, 'Dashed line: identity.'


# ---------------------------------------------------------------------------------------------- Figure 4
@conv('panel_classes')
def c_classes(P):
    rows = OrderedDict()
    for c in P.get('errorbar'):
        r = c.ctx['r']
        d = rows.setdefault(r['class'], dict(mechanism_class=r['class'], drugs=int(r['drugs']),
                                             n_reproducible=r['n_reproducible'],
                                             row_label=tick_at(c.ax, 'y', c.ctx['k'])))
        if fmt_marker(c) == 'D':
            a_ = c.ctx['a_']
            d.update(activity_adjusted_mean=num(c.args[0])[0], adjusted_ci_lo=a_.adj_ci95_lo, adjusted_ci_hi=a_.adj_ci95_hi)
        else:
            d.update(median_selective_agreement=num(c.args[0])[0], median_ci_lo=r['median_ci95_lo'],
                     median_ci_hi=r['median_ci95_hi'], share_reproducible=r['frac_reproducible'])
    cols = OrderedDict(mechanism_class='manually curated mechanism class (≥5 drugs)', drugs='drugs in the class',
                       n_reproducible='reproducible drugs (agreement ≥ 0.3)', row_label='row label',
                       median_selective_agreement='median cross-screen selective agreement (filled circle)',
                       median_ci_lo='95% drug-bootstrap interval, lower', median_ci_hi='95% drug-bootstrap interval, upper',
                       activity_adjusted_mean='class mean adjusted for drug activity (open diamond; logit-quadratic '
                                              'specification)',
                       adjusted_ci_lo='95% interval, lower', adjusted_ci_hi='95% interval, upper',
                       share_reproducible='share of the class reproducible (percentage at right)')
    df = T(list(rows.values()), cols)
    chk = [check(df, 'median_selective_agreement', 'results/revision_transfer/class_summary_curated.csv',
                 [('mechanism_class', 'class')], 'median_rho_sel', up_filter=lambda U: U[U.scheme == 'curated']),
           check(df, 'activity_adjusted_mean', 'results/revision_transfer/activity_regression.csv',
                 [('mechanism_class', 'class')], 'adjusted_mean',
                 up_filter=lambda U: U[(U.scheme == 'curated') & (U.spec == 'logit_quadratic')])]
    return df, cols, chk, ''


@conv('panel_c2')
def c_c2(P):
    leg = P.legend
    rows = []
    for c in P.get('errorbar'):
        if fmt_marker(c) == 'o':
            r = c.ctx['r']
            rows.append(dict(drug_set=legend_label(leg, col=c.kw.get('color')), drug_set_code=c.ctx['cls'],
                             predictor=c.ctx['lab'], estimand_model=c.ctx['m'], c2=num(c.args[0])[0], ci_lo=r.ci_lo,
                             ci_hi=r.ci_hi, n_drugs=r.n_drugs))
        else:
            B, lo, hi = c.ctx['B'], c.ctx['lo'], c.ctx['hi']
            rows.append(dict(drug_set=f'{int(B.n_drugs)} drugs with a canonical biomarker (core)',
                             drug_set_code='core, primary biomarker per drug', predictor=c.ctx['lab'],
                             estimand_model=c.ctx['c'], c2=num(c.args[0])[0], ci_lo=B[lo] if lo else np.nan,
                             ci_hi=B[hi] if hi else np.nan, n_drugs=B.n_drugs))
    cols = OrderedDict(drug_set='drugs pooled (legend; lower block: drugs with a canonical biomarker)',
                       drug_set_code='class / set in the upstream file', predictor='predictor (row label)',
                       estimand_model='model (49_c2_estimands.csv) or column (52_biomarker_c2_pooled.csv)',
                       c2='share of reproducible selective variance captured (c², pooled over drugs)',
                       ci_lo='95% drug-cluster interval, lower', ci_hi='95% drug-cluster interval, upper',
                       n_drugs='drugs pooled')
    df = T(rows, cols)
    m = df.drug_set_code.isin(['all', 'reproducible'])
    chk = [check(df, 'c2', 'results/revision_ceiling/49_c2_estimands.csv',
                 [('estimand_model', 'model'), ('drug_set_code', 'class')], 'estimate', mask=m,
                 up_filter=lambda U: U[(U.kappa == 'none') & (U.aggregation == 'perdrug_romeans_bc')])]
    B = upstream('results/revision_ceiling/52_biomarker_c2_pooled.csv').set_index('set').loc['core, primary biomarker per drug']
    sub = df[~m]
    chk.append(check_values('c2 (canonical-biomarker block)', sub.c2, [B[c] for c in sub.estimand_model],
                            'results/revision_ceiling/52_biomarker_c2_pooled.csv', 'c2_*_pooled (set: core, primary)'))
    return df, cols, chk, ''


@conv('panel_biomarker')
def c_biomarker(P):
    named = {c.ctx['r']['drug_name']: plain(c.args[0]) for c in P.get('annotate')}
    rows = []
    for c in P.get('scatter'):
        D = c.ctx['D']
        for i, x, y in zip(sidx(c.args[0]), num(c.args[0]), num(c.args[1])):
            r = D.loc[i]
            rows.append(dict(drug=r.drug_name, compound=r.compound, biomarker=r.biomarker, tier=r.tier,
                             c2_kernel_ridge=x, c2_biomarker=y, label_in_figure=named.get(r.drug_name, '')))
    for c in P.get('errorbar'):
        Pj = c.ctx['P']
        rows.append(dict(drug='pooled over the 29 core drugs', tier='core (pooled)', c2_kernel_ridge=num(c.args[0])[0],
                         c2_kernel_ridge_ci_lo=Pj['ridge_ci_lo'], c2_kernel_ridge_ci_hi=Pj['ridge_ci_hi'],
                         c2_biomarker=num(c.args[1])[0], c2_biomarker_ci_lo=Pj['ci_lo'], c2_biomarker_ci_hi=Pj['ci_hi'],
                         label_in_figure='square'))
    cols = OrderedDict(drug='drug', compound='harmonized compound id', biomarker='pre-specified canonical biomarker',
                       tier='core (aqua) or extended (light blue) biomarker; pooled square',
                       c2_kernel_ridge='c² of the selective-target kernel ridge (x)',
                       c2_biomarker='c² of the canonical biomarker (y)',
                       label_in_figure='annotation text (labelled drugs) or square (pooled estimate)',
                       c2_kernel_ridge_ci_lo='pooled square: 95% drug-cluster interval, lower',
                       c2_kernel_ridge_ci_hi='95% interval, upper',
                       c2_biomarker_ci_lo='pooled square: 95% drug-cluster interval, lower',
                       c2_biomarker_ci_hi='95% interval, upper')
    df = T(rows, cols)
    m = df.compound.notna()
    k = [('compound', 'compound'), ('biomarker', 'biomarker')]
    U = 'results/revision_ceiling/52_biomarker_c2_perdrug.csv'
    pj = upstream('results/revision_ceiling/52_summary.json')['pooled'][0]
    sq = df[~m]
    chk = [check(df, 'c2_kernel_ridge', U, k, 'c2_ridge_sel', mask=m),
           check(df, 'c2_biomarker', U, k, 'c2_biomarker', mask=m),
           check_values('c2_biomarker (pooled square)', sq.c2_biomarker, [pj['c2_biomarker_pooled']],
                        'results/revision_ceiling/52_summary.json', 'pooled[0].c2_biomarker_pooled')]
    return df, cols, chk, 'Line: identity.'


@conv('panel_crossdataset')
def c_crossdataset(P):
    leg = P.legend
    rows = []
    for c in P.get('scatter'):
        for m, y in zip(sidx(c.args[1]), num(c.args[1])):
            rows.append(dict(target=c.ctx['t'], pair=c.ctx['pair'], series='published models', model=m,
                             statistic='selB_c', median_selective_both=y))
    for c in P.get('errorbar'):
        r = c.ctx['r']
        mk = fmt_marker(c)
        lab = legend_label(leg, marker=mk, col=c.kw.get('color'))
        rows.append(dict(target=c.ctx['t'], pair=c.ctx['pair'], series=lab, model=c.ctx['m'],
                         statistic='adj_c' if mk == 'D' else 'selB_c', median_selective_both=num(c.args[1])[0],
                         ci_lo=r.median_lo, ci_hi=r.median_hi))
    for c in P.get('plot'):
        rows.append(dict(target=c.ctx['t'], pair=c.ctx['pair'], series="CTRPv2's own measurement (one replicate)",
                         model='ref_CTRPv2_measured', statistic='selB_c', median_selective_both=num(c.args[1])[0]))
    for c in P.get('add_patch'):
        rows.append(dict(target=c.ctx['t'], pair=c.ctx['pair'], series='drug-agnostic floor (shaded)',
                         model='max over drug-agnostic predictors', statistic='selBOTH (leakage_floors.csv)',
                         median_selective_both=c.ctx['fl']))
    cols = OrderedDict(target='dataset the CTRPv2-trained predictions are scored in',
                       pair='training -> target, as in the upstream files',
                       series='published models (grey), EnsembleMF, kernel ridge, own axis removed (open diamonds), '
                              "CTRPv2's measurement (dashed), drug-agnostic floor (shading)",
                       model='model', statistic='selB_c: selective-both; adj_c: with the model\'s own axis removed; '
                                                'floor: maximum median selective-both of the drug-agnostic predictors',
                       median_selective_both='median over drugs of the selective-both correlation in the target',
                       ci_lo='95% two-way bootstrap interval, lower', ci_hi='95% interval, upper')
    df = T(rows, cols)
    Mm = 'results/revision_crosslab/model_medians.csv'
    k = [('pair', 'pair'), ('model', 'model'), ('statistic', 'stat')]
    f = lambda U: U[U.subset == 'all']
    pub = df.series == 'published models'
    chk = [check(df, 'median_selective_both', Mm, k, 'median', mask=pub, up_filter=f),
           check(df, 'median_selective_both', Mm, k, 'median', mask=~pub & df.statistic.isin(['selB_c', 'adj_c']),
                 up_filter=f, note='EnsembleMF, kernel ridge and the CTRPv2 measurement')]
    return df, cols, chk, ''


# ---------------------------------------------------------------------------------------------- Figure 5
@conv('panel_estimates')
def c_estimates(P):
    rows = OrderedDict()
    thr = {}
    for c in P.calls:
        if 'h' not in c.ctx or 's' not in c.ctx or c.method not in ('errorbar', 'scatter', 'plot', 'axvline'):
            continue
        h, s = c.ctx['h'], c.ctx['s']
        if c.method == 'axvline':
            if c.kw.get('color') == M.ORANGE:
                thr[h] = (num(c.args[0])[0], c.ctx['kind'])
            continue
        d = rows.setdefault((h, s), dict(hypothesis=h, quantity=c.ctx['q'], panel_title=plain(c.ctx['title']),
                                         screen=c.ctx['lab'], screen_code=s))
        if c.method == 'errorbar' and fmt_marker(c) == '^':
            d.update(drug_agnostic_floor=num(c.args[0])[0], floor_ci_lo=c.ctx['flo'], floor_ci_hi=c.ctx['fhi'])
        elif c.method == 'errorbar':
            ok = c.ctx['ok']
            d.update(estimate=num(c.args[0])[0], ci_lo=c.ctx['lo'], ci_hi=c.ctx['hi'],
                     registered_test='not estimable' if ok is None else ('met' if ok else 'not met'))
        elif c.method == 'scatter' and c.kw.get('marker') == 'D':
            d['discovery_value'] = num(c.args[0])[0]
        elif c.method == 'plot':
            d['discovery_range_lo'], d['discovery_range_hi'] = num(c.args[0])
        elif c.method == 'scatter' and c.kw.get('marker') == 'x':
            d['gdsc2_rehearsal'] = num(c.args[0])[0]
    for d in rows.values():
        if d['hypothesis'] in thr:
            d['registered_threshold'] = thr[d['hypothesis']][0]
            d['threshold_rule'] = {'lt': 'test met if estimate < threshold',
                                   'ge': 'test met if estimate ≥ threshold'}[thr[d['hypothesis']][1]]
    cols = OrderedDict(hypothesis='registered hypothesis (sub-panel)', quantity='quantity as named in the upstream file',
                       panel_title='sub-panel title', screen='untouched confirmation screen (row)',
                       screen_code='screen code in the upstream file',
                       estimate='estimate (median or ratio); filled: registered test met, open: not met, grey: not estimable',
                       ci_lo='95% two-way (drug × cell line) bootstrap interval, lower', ci_hi='95% interval, upper',
                       registered_test='registered pass/fail (corrected H4 flag, as in Supplementary Note 2)',
                       discovery_value='discovery value (diamond)',
                       discovery_range_lo='range over discovery screens (line), lower',
                       discovery_range_hi='range over discovery screens, upper',
                       gdsc2_rehearsal='GDSC2 rehearsal (cross; the larger value where two are recorded)',
                       drug_agnostic_floor='H5a: frozen drug-agnostic discovery axis scored the same way (triangle)',
                       floor_ci_lo='95% interval, lower', floor_ci_hi='95% interval, upper',
                       registered_threshold='registered threshold (orange line)', threshold_rule='direction of the test')
    df = T(list(rows.values()), cols)
    E = 'results/revision_transfer/confirmation_estimates.csv'
    chk = [check(df, 'estimate', E, [('screen_code', 'screen'), ('hypothesis', 'hypothesis'), ('quantity', 'quantity')],
                 'estimate'),
           check(df, 'ci_lo', E, [('screen_code', 'screen'), ('hypothesis', 'hypothesis'), ('quantity', 'quantity')],
                 'ci95_two_way_boot', transform=lambda v: ci_pair(v)[0], note='parsed from the interval string')]
    return df, cols, chk, ''


@conv('panel_patient')
def c_patient(P):
    pv = {}
    for c in P.get('text'):
        if 'labs' in c.ctx and 'k' in c.ctx and c.ctx['k'] < len(c.ctx['labs']) and 'P' in str(c.args[2]):
            pv[c.ctx['labs'][c.ctx['k']]] = plain(str(c.args[2]).replace('\n', '; '))
    rows = []
    for c in P.get('errorbar'):
        x = c.ctx
        rows.append(dict(cohort_unit=plain(x['lab']), drug_class=x['cls'], pooled_r=num(c.args[0])[0], ci95_lo=x['lo'],
                         ci95_hi=x['hi'], pi95_lo=x['plo'], pi95_hi=x['phi'], p_values=pv.get(x['lab'], '')))
    cols = OrderedDict(cohort_unit='cohort and unit of analysis (row)',
                       drug_class='discovery reproducibility class (blue: reproducible; grey: non-reproducible)',
                       pooled_r='pooled selective transfer r (random-effects meta-analysis over drugs)',
                       ci95_lo='Hartung–Knapp 95% interval, lower', ci95_hi='Hartung–Knapp 95% interval, upper',
                       pi95_lo='95% prediction interval for a new drug, lower (thin line; not drawn for organoids)',
                       pi95_hi='95% prediction interval, upper',
                       p_values='P values printed for the row (registered tests; see caption)')
    df = T(rows, cols)
    chk = []
    B = upstream('results/revision_transfer/beataml_meta_analysis.csv')
    Px = upstream('results/revision_transfer/pdx_meta_analysis.csv')
    Pm = upstream('results/revision_transfer/pdm_meta_analysis.csv')
    look = {'BeatAML, samples (registered)': lambda c: B[(B.analysis == 'registered sample level') & (B['class'] == c)].pooled_r,
            'BeatAML, patients': lambda c: B[(B.analysis == 'patient level') & (B['class'] == c)].pooled_r,
            'PDX, models (registered)': lambda c: Px[(Px.analysis == 'registered model level') & (Px.grouping == 'model')
                                                     & (Px['class'] == c)].pooled_r,
            'PDX, tumours': lambda c: Px[(Px.analysis == 'T1 tumour level') & (Px.grouping == 'G1') & (Px['class'] == c)].pooled_r,
            'Organoids, patients (corrected orientation)': lambda c: Pm[Pm.analysis.str.startswith('corrected')
                                                                        & (Pm.scope == 'all') & (Pm['class'] == c)].re_pooled_r}
    exp = [look[r.cohort_unit](r.drug_class).iloc[0] if r.cohort_unit in look else np.nan for r in df.itertuples()]
    chk.append(check_values('pooled_r', df.pooled_r, exp, 'results/revision_transfer/{beataml,pdx,pdm}_meta_analysis.csv',
                            'pooled_r / re_pooled_r', note='row-by-row lookup by analysis and class'))
    return df, cols, chk, ''


@conv('panel_patient_transfer')
def c_patient_transfer(P):
    labels = {}
    for c in P.get('text'):
        if 'r' in c.ctx and 'nm' in c.ctx:
            labels[c.ctx['r'].compound] = plain(c.ctx['nm'])
    rows = []
    folder = None
    for c in P.get('scatter'):
        v, comp, gen = c.ctx['v'], c.ctx['comp'], c.ctx['gen']
        folder = c.ctx['folder']
        for i, y in zip(sidx(c.args[1]), num(c.args[1])):
            r = v.loc[i]
            row = dict(dataset=folder, drug_class=c.ctx['cls'], compound=r.compound,
                       drug=gen.get(r.compound, str(comp.get(r.compound, ''))))
            if 'dataset' in v.columns:
                row['organoid_set'] = r.dataset
            row.update(n_samples=r.n, selective_transfer=y, label_in_figure=labels.get(r.compound, ''), element='drug')
            rows.append(row)
    for c in P.get('plot'):
        if c.ctx.get('cls') and len(num(c.args[1])) == 2 and num(c.args[1])[0] == num(c.args[1])[1] and 'xi' not in c.ctx:
            rows.append(dict(dataset=c.ctx['folder'], drug_class=c.ctx['cls'], selective_transfer=num(c.args[1])[0],
                             element='class median (bar)'))
    cols = OrderedDict(dataset='results folder of the confirmation (BeatAML, PDX-derived cells, organoids)',
                       drug_class='discovery reproducibility class (column)', compound='harmonized compound id',
                       drug='drug name', organoid_set='organoid dataset (liver or sarcoma; one unit per drug and dataset)',
                       n_samples='samples (or units) scored for the drug (in parentheses in labels when < 100)',
                       selective_transfer='selective transfer per drug–dataset unit (correlation of the frozen '
                                          'prediction with the selective response)',
                       label_in_figure='label printed in the panel (†: MEK inhibitor)', element='drug or class median')
    df = T(rows, cols)
    chk = []
    if folder:
        keys = [('compound', 'compound')] + ([('organoid_set', 'dataset')] if 'organoid_set' in df else [])
        chk.append(check(df, 'selective_transfer', f'results/{folder}/scores.csv', keys, 'sel',
                         mask=df.element == 'drug', up_filter=lambda U: U[U.family == 'sel']))
    return df, cols, chk, ''


# ---------------------------------------------------------------------------------------------- supplementary
@conv('sfigure1')
def c_s_growth(P):
    rows = []
    for c in P.get('scatter'):
        s = c.ctx['s']
        for i, x, y in zip(sidx(c.args[0]), num(c.args[0]), num(c.args[1])):
            rows.append(dict(screen=F.LABEL.get(s, s), screen_code=s, cell_line=i, growth_rate=x, axis=y))
    cols = OrderedDict(screen='screen (sub-panel)', screen_code='column in the upstream file',
                       cell_line='DepMap cell-line id', growth_rate="DepMap CRISPR-inferred growth rate (x)",
                       axis="the screen's all-drug general-sensitivity axis (mean z; y)")
    df = T(rows, cols)
    A = READ_CSV(PROJECT / 'results/fitness_axis/fitness_axes.csv', index_col=0)
    G = READ_CSV(PROJECT / 'data/processed/features/growth_rate.csv', index_col=0).growth_rate
    chk = [check_values('axis', df.axis, [A.loc[r.cell_line, r.screen_code] for r in df.itertuples()],
                        'results/fitness_axis/fitness_axes.csv', '<screen> column by cell line'),
           check_values('growth_rate', df.growth_rate, [G.loc[r.cell_line] for r in df.itertuples()],
                        'data/processed/features/growth_rate.csv', 'growth_rate by cell line')]
    return df, cols, chk, ''


@conv('sfigure7')
def c_s_normalized(P):
    c = P.get('barh')[0]
    R = c.ctx['R']
    rows = []
    for y, i, w in zip(num(c.args[0]), sidx(c.args[1]), num(c.args[1])):
        r = R.loc[i]
        rows.append(dict(model=r.model, model_label=tick_at(c.ax, 'y', y), model_class=r['class'], normalized_pearson=w,
                         split_sd=r.split_sd))
    cols = OrderedDict(model='model', model_label='row label', model_class='class (blue: leaderboard; grey: naive; '
                                                                          'green: this study)',
                       normalized_pearson="DrEval normalized Pearson, mean over 10 splits (bar)",
                       split_sd='s.d. over splits (error bar)')
    df = T(rows, cols)
    chk = [check(df, 'normalized_pearson', 'results/dreval/normalized_metric.csv', [('model', 'model')], 'normalized_pearson')]
    return df, cols, chk, ''


@conv('sfigure6')
def c_s_modes(P):
    rows = OrderedDict()
    for c in P.get('scatter'):
        d = c.ctx['d']
        col = c.ctx['a'] if 'alpha' in c.kw else c.ctx['b']
        for i, y in zip(sidx(c.args[1]), num(c.args[1])):
            alg = d.loc[i, 'algorithm']
            r = rows.setdefault((c.ctx['mode'], c.ctx['t'], alg), dict(test_mode=c.ctx['mode'], mode_label=c.ctx['lab'],
                                                                       target=c.ctx['t'], model=alg))
            r[col] = y
    cols = OrderedDict(test_mode='DrEval test mode (sub-panel)', mode_label='test mode description',
                       target='dataset the CTRPv2-trained predictions are scored in (x group)', model='DrEval model',
                       raw_within='median per-drug correlation scored in CTRPv2 (blue, faded)',
                       raw_cross='median per-drug correlation scored in the target (blue)',
                       sel_within='median selective correlation in CTRPv2 (orange, faded)',
                       sel_cross='median selective correlation in the target (orange; target axis only)')
    df = T(list(rows.values()), cols)
    k = [('model', 'algorithm'), ('test_mode', 'mode'), ('target', 'target')]
    chk = [check(df, c, 'results/dreval_crossstudy/summary.csv', k, c) for c in ('raw_cross', 'sel_cross')]
    return df, cols, chk, 'Bars: medians over models (not exported; computable from the rows).'


@conv('sfigure_example:a')
def c_s_example_a(P):
    rows = []
    for c in P.get('scatter'):
        for i, x, y in zip(sidx(c.args[0]), num(c.args[0]), num(c.args[1])):
            rows.append(dict(cell_line=i, group=c.ctx['k'], axis_g_sd=x, trametinib_z=y))
    pl = P.get('plot')
    b0, b1 = (pl[0].ctx['b0'], pl[0].ctx['b1']) if pl else (np.nan, np.nan)
    cols = OrderedDict(cell_line='DepMap cell-line id', group='BRAF or NRAS/KRAS hotspot status (colour)',
                       axis_g_sd='general-sensitivity axis g of CTRP, standardized (x)',
                       trametinib_z='trametinib sensitivity z in CTRP (y)')
    df = T(rows, cols)
    z = READ_PARQUET(PROJECT / 'results/decomposition/z_CTRP.parquet').loc['DPC-006587']
    chk = [check_values('trametinib_z', df.trametinib_z, [z[c] for c in df.cell_line],
                        'results/decomposition/z_CTRP.parquet', 'row DPC-006587 (trametinib), by cell line')]
    return df, cols, chk, f'Line: least-squares fit z = {b0:.6g} + {b1:.6g} g.'


@conv('sfigure_example:b')
def c_s_example_b(P):
    rows = []
    for c in P.get('scatter'):
        df0, gname, col = c.ctx['df'], c.ctx['gname'], c.ctx['col']
        cells = df0.index[df0.grp == gname]
        for i, y in zip(cells, num(c.args[1])):
            rows.append(dict(group=gname, measure='profile z' if col == 'z' else 'selective residual', cell_line=i,
                             value=y, element='cell line'))
    for c in P.get('plot'):
        rows.append(dict(group=c.ctx['gname'], measure='profile z' if c.ctx['col'] == 'z' else 'selective residual',
                         value=num(c.args[1])[0], element='median (bar)'))
    cols = OrderedDict(group='mutation group (x)', measure='left of each pair: profile z (grey); right: selective '
                                                           'residual (coloured)',
                       cell_line='DepMap cell-line id', value='trametinib z or selective residual in CTRP',
                       element='cell line or group median')
    df = T(rows, cols)
    r = READ_PARQUET(PROJECT / 'results/decomposition/resid_CTRP.parquet').loc['DPC-006587']
    sub = df[(df.element == 'cell line') & (df.measure == 'selective residual')]
    chk = [check_values('value (selective residual)', sub.value, [r[c] for c in sub.cell_line],
                        'results/decomposition/resid_CTRP.parquet', 'row DPC-006587 (trametinib), by cell line')]
    return df, cols, chk, ''


@conv('sfigure_example:c')
def c_s_example_c(P):
    rows = []
    for c in P.get('plot'):
        for x, y in zip(num(c.args[0]), num(c.args[1])):
            rows.append(dict(curve=plain(c.kw.get('label')), per_drug_correlation=x, mean_selected_response_sd=y))
    for c in P.get('axvline'):
        rows.append(dict(curve='reference: ' + plain(c.ctx['lab']), per_drug_correlation=num(c.args[0])[0]))
    cols = OrderedDict(curve='selection rule (line) or reference value (vertical line)',
                       per_drug_correlation='per-drug correlation of the prediction (x)',
                       mean_selected_response_sd='mean response of the selected lines, s.d. above average (bivariate '
                                                 'normal: r × φ(c)/q for the top fraction q)')
    df = T(rows, cols)
    from scipy.stats import norm
    chk = []
    for q in (0.1, 0.25):
        sub = df[df.curve.str.startswith(f'top {int(q * 100)}%')]
        cq = norm.ppf(1 - q)
        chk.append(check_values(f'mean_selected_response_sd (top {int(q * 100)}%)', sub.mean_selected_response_sd,
                                sub.per_drug_correlation * (norm.pdf(cq) / q), 'analytic (no upstream file)',
                                'r · φ(Φ⁻¹(1-q)) / q', note='recomputation of the analytic curve'))
    return df, cols, chk, 'Analytic panel; reference lines are values quoted in the text.'


@conv('sfigure3:a')
def c_s_dreval_a(P):
    c = P.get('scatter')[0]
    rows = [dict(drug=i, EnsembleMF_selective_r=x, kernel_ridge_selective_r=y)
            for i, x, y in zip(sidx(c.args[0]), num(c.args[0]), num(c.args[1]))]
    cols = OrderedDict(drug='DrEval drug id', EnsembleMF_selective_r='per-drug selective correlation of EnsembleMF (x)',
                       kernel_ridge_selective_r='per-drug selective correlation of the selective-target kernel ridge, '
                                                'landmark genes (y)')
    df = T(rows, cols)
    U = READ_CSV(PROJECT / 'results/dreval/per_drug.csv.gz', dtype={'drug': str})
    W = U.pivot_table(index='drug', columns='model', values='selective')
    chk = [check_values('kernel_ridge_selective_r', df.kernel_ridge_selective_r, W.loc[df.drug, 'ridge_sel_lm'],
                        'results/dreval/per_drug.csv.gz', 'selective (model ridge_sel_lm, by drug)')]
    return df, cols, chk, 'Line: identity.'


@conv('sfigure3:b')
def c_s_dreval_b(P):
    rows = OrderedDict()
    for c in P.get('scatter'):
        Td = c.ctx['T']
        for x, i, y in zip(num(c.args[0]), sidx(c.args[1]), num(c.args[1])):
            r = rows.setdefault(i, dict(model=Td.loc[i, 'model'], model_label=tick_at(c.ax, 'x', x)))
            r['selective_all_drugs' if c.kw.get('label') == 'all drugs' else 'selective_reproducible_drugs'] = y
    cols = OrderedDict(model='leaderboard model', model_label='tick label',
                       selective_all_drugs='median selective correlation over all drugs (grey)',
                       selective_reproducible_drugs='median over the drugs reproducible in discovery (blue)')
    df = T(list(rows.values()), cols)
    chk = [check(df, 'selective_reproducible_drugs', 'results/dreval/model_summary.csv', [('model', 'model')],
                 'reproducible_selective')]
    return df, cols, chk, ''


@conv('sfigure3:c')
def c_s_dreval_c(P):
    rows = []
    for c in P.get('bar'):
        raw = color(c.kw.get('color')) == color('#86b6ef')
        rows.append(dict(input_genes=plain(c.ctx['lab']), training_target='raw' if raw else 'selective',
                         model=('ridge_raw' if raw else 'ridge_sel') + c.ctx['sfx'], median_selective_r=num(c.args[1])[0]))
    for c in P.get('axhline'):
        rows.append(dict(input_genes='', training_target='', model='EnsembleMF (best published; dashed line)',
                         median_selective_r=num(c.args[0])[0]))
    cols = OrderedDict(input_genes='input gene set (x group)', training_target='trained on the raw (light blue) or '
                                                                              'selective (green) target',
                       model='model', median_selective_r='median selective correlation over drugs')
    df = T(rows, cols)
    df['key'] = df.model.str.replace(r' \(.*', '', regex=True)
    chk = [check(df, 'median_selective_r', 'results/dreval/model_summary.csv', [('key', 'model')], 'selective')]
    return df.drop(columns='key'), cols, chk, ''


@conv('panel_own_models')
def c_s_own_models(P):
    rows = []
    models = None
    for c in P.get('bar'):
        models = c.ctx['models']
        for x, h in zip(num(c.args[0]), num(c.args[1])):
            k = int(round(x))
            rows.append(dict(model=models[k][0], model_label=plain(models[k][1]), series=c.kw.get('label'),
                             median_r=h))
    cols = OrderedDict(model='cell-blind predictor', model_label='tick label',
                       series='bar (legend): within screen per-drug (grey); other screens per-drug (blue); other '
                              'screens selective (orange)',
                       median_r='median over drug × train/evaluation-screen combinations')
    df = T(rows, cols)
    s = upstream('results/models/score_summary.csv')
    w = df[df.series.str.startswith('within')]
    exp = s[s['train'] == s['eval']].groupby('model').raw.median()
    chk = [check_values('median_r (within screen)', w.median_r, exp.loc[w.model],
                        'results/models/score_summary.csv', 'raw (median where train == eval)',
                        note='recomputation: median as in the panel code')]
    return df, cols, chk, ''


@conv('panel_global_models')
def c_s_global(P):
    rows = []
    for c in P.get('bar'):
        tab, colname = c.ctx['tab'], c.ctx['col']
        order = c.ctx['order']
        for (k, lab), h in zip(order, num(c.args[1])):
            rows.append(dict(split=k, split_label=lab, model='trees' if c.kw.get('label', '').startswith('trees')
                             else 'neural net', metric='per-drug' if colname == 'raw' else 'selective',
                             median_r=h, global_pearson_r=tab.loc[k, 'glob']))
    for c in P.get('plot'):
        rows.append(dict(split='random_pairs', split_label='random pairs', model='drug + cell means (no features)',
                         metric='per-drug', median_r=num(c.args[1])[0]))
    cols = OrderedDict(split='split design', split_label='tick label', model='gradient-boosted trees, two-tower neural '
                                                                            'network, or drug + cell means (black bar)',
                       metric='per-drug (raw) or selective correlation', median_r='median over drugs and screens',
                       global_pearson_r='global Pearson r over all pairs (in the tick label)')
    df = T(rows, cols)
    gsc = upstream('results/global_model/scores.csv')
    gg = gsc[gsc.model == 'global_gbt'].copy()
    gg['split2'] = gg.split.where(~gg.split.str.startswith('cross_from'), 'cross_screen')
    exp = gg.groupby('split2').per_drug_raw.median()
    sub = df[(df.model == 'trees') & (df.metric == 'per-drug')]
    chk = [check_values('median_r (trees, per-drug)', sub.median_r, exp.loc[sub.split], 'results/global_model/scores.csv',
                        'per_drug_raw (median by split)', note='recomputation: median as in the panel code')]
    return df, cols, chk, ''


@conv('panel_dreval_collapse')
def c_s_collapse(P):
    rows = []
    for c in P.get('scatter'):
        r = c.ctx['r']
        rows.append(dict(model=r.model, model_class=r.klass, drug_agnostic=c.kw.get('marker') == '^',
                         prediction_axis_share=num(c.args[0])[0], median_selective_r=num(c.args[1])[0]))
    for c in P.get('axvline'):
        rows.append(dict(model='observed responses (dashed line)', prediction_axis_share=num(c.args[0])[0]))
    cols = OrderedDict(model='leaderboard model', model_class='class (colour)', drug_agnostic='triangle',
                       prediction_axis_share="shared-axis share of predicted profiles (median R² on the other drugs' "
                                             "mean prediction; x)",
                       median_selective_r='median selective correlation over drugs (y)')
    df = T(rows, cols)
    chk = [check(df, 'prediction_axis_share', 'results/dreval/model_summary.csv', [('model', 'model')],
                 'prediction_fitness_r2', mask=df.model_class.notna())]
    return df, cols, chk, ''


@conv('sfigure5:a')
def c_s_voi_a(P):
    rows = []
    for c in P.get('errorbar'):
        x = c.ctx
        rows.append(dict(features=x['lab'], features_code=x['f'], evaluated_in=F.LABEL.get(x['sname'], x['sname']),
                         eval_code=x['sname'], map_r=num(c.args[0])[0], ci_lo=x['lo'], ci_hi=x['hi']))
    cols = OrderedDict(features='response features (row)', features_code='features in the upstream file',
                       evaluated_in='discovery screen whose selective residual is predicted (colour)',
                       eval_code='screen code', map_r='drug-specific correlation with the selective residual '
                                                      '(drug-blind ridge map)',
                       ci_lo='95% drug-bootstrap interval, lower', ci_hi='95% interval, upper')
    df = T(rows, cols)
    k = [('features_code', 'features'), ('eval_code', 'eval')]
    chk = [check(df, 'map_r', 'results/value_of_information/l1000_two_way.csv', k, 'map_r',
                 mask=df.features_code.isin(set(upstream('results/value_of_information/l1000_two_way.csv').features))
                 & ~df.features_code.isin(['nearest_chemistry', 'same_target']))]
    return df, cols, chk, 'Shaded rows: retrieval baselines (results/value_of_information/retrieval_baselines.csv).'


@conv('sfigure5:b')
def c_s_voi_b(P):
    rows = []
    for c in P.get('plot'):
        ks = c.ctx['ks']
        for x, y in zip(num(c.args[0]), num(c.args[1])):
            k = int(x)
            rows.append(dict(evaluated_in=F.LABEL.get(c.ctx['sname'], c.ctx['sname']), eval_code=c.ctx['sname'],
                             variant=tick_at(c.ax, 'x', k), features_code=ks[k], map_r=y))
    cols = OrderedDict(evaluated_in='discovery screen (line colour)', eval_code='screen code',
                       variant='measured response with leading shared principal components removed (x)',
                       features_code='features in the upstream file', map_r='correlation with the selective residual')
    df = T(rows, cols)
    chk = [check(df, 'map_r', 'results/value_of_information/retrieval_baselines.csv',
                 [('features_code', 'features'), ('eval_code', 'eval')], 'map_r')]
    return df, cols, chk, ''


@conv('sfigure5:c')
def c_s_voi_c(P):
    rows = []
    for c in P.get('scatter'):
        d = c.ctx['d']
        for i, x, y in zip(sidx(c.args[0]), num(c.args[0]), num(c.args[1])):
            rows.append(dict(evaluated_in=F.LABEL.get(c.ctx['sname'], c.ctx['sname']), eval_code=c.ctx['sname'],
                             basal_trained_in=d.loc[i, 'basal_train'], basal_r=x, combined_r=y,
                             measured_given_basal=d.loc[i, 'measured_given_basal']))
    cols = OrderedDict(evaluated_in='discovery screen evaluated (colour)', eval_code='screen code',
                       basal_trained_in='screen the basal-omics model was trained in',
                       basal_r='basal-omics model (x)', combined_r='basal model + measured response (y)',
                       measured_given_basal='partial r of the measured response given the basal model (median in text)')
    df = T(rows, cols)
    chk = [check(df, 'combined_r', 'results/value_of_information/l1000_incremental.csv',
                 [('eval_code', 'eval'), ('basal_trained_in', 'basal_train')], 'combined_r')]
    return df, cols, chk, 'Line: identity.'


@conv('sfigure2')
def c_s_tahoe(P):
    rows = []
    for c in P.get('scatter'):
        rep = c.kw.get('marker') == 's'
        s = c.ctx['s']
        rows.append(dict(features='viability in another screen (replicate; median over screens)' if rep
                         else plain(c.ctx['lab']), features_code='replicate_*' if rep else c.ctx['f'],
                         evaluated_in=F.LABEL.get(s, s), eval_code=s, pooled_within_drug_r=num(c.args[0])[0]))
    cols = OrderedDict(features='Tahoe-100M response features (row); squares: another screen\'s selective residual',
                       features_code='features in the upstream file', evaluated_in='discovery screen (colour)',
                       eval_code='screen code',
                       pooled_within_drug_r='pooled within-drug correlation with the selective residual')
    df = T(rows, cols)
    chk = [check(df, 'pooled_within_drug_r', 'results/value_of_information/tahoe_pooled.csv',
                 [('features_code', 'features'), ('eval_code', 'eval')], 'ridge_map_vs_selective',
                 mask=df.features_code != 'replicate_*')]
    t = upstream('results/value_of_information/tahoe_pooled.csv')
    rep = t[t.features.str.startswith('replicate_')].groupby('eval').ridge_map_vs_selective.median()
    sub = df[df.features_code == 'replicate_*']
    chk.append(check_values('pooled_within_drug_r (replicate squares)', sub.pooled_within_drug_r, rep.loc[sub.eval_code],
                            'results/value_of_information/tahoe_pooled.csv', 'ridge_map_vs_selective (median of '
                            'replicate_* by eval)', note='recomputation: median as in the panel code'))
    return df, cols, chk, ''


# ---------------------------------------------------------------------------------------------- Figure 1 (headline)
EX1 = 'results/revision_round4/fig1_example'


@conv('panel_example_profile')
def c_example_profile(P):
    c = P.get('scatter')[0]
    C, Tt = c.ctx['C'], c.ctx['T']
    drug = Tt.name.iloc[0]
    rows = [dict(drug=drug, cell_line=C.loc[i, 'cell'], axis_g=x, profile_z=y)
            for i, x, y in zip(sidx(c.args[0]), num(c.args[0]), num(c.args[1]))]
    pl = P.get('plot')
    b0, b1 = (pl[0].ctx['b0'], pl[0].ctx['b1']) if pl else (np.nan, np.nan)
    cols = OrderedDict(drug='example drug', cell_line='cell line (Cellosaurus id)',
                       axis_g='general-sensitivity axis g: leave-drug-out mean of z over the other drugs (x)',
                       profile_z="the drug's profile z, rank-normal (y)")
    df = T(rows, cols)
    k = [('cell_line', 'cell')]
    chk = [check(df, 'profile_z', f'{EX1}/cells.csv', k, 'z'), check(df, 'axis_g', f'{EX1}/cells.csv', k, 'g')]
    note = (f'{drug} (DrEval drug {Tt.drug.iloc[0]}), {int(Tt.cells.iloc[0])} cell lines; r_zg = {Tt.r_zg.iloc[0]:.6g}. '
            f'Line: least-squares fit z = {b0:.6g} + {b1:.6g} g.')
    return df, cols, chk, note


@conv('panel_example_terms')
def c_example_terms(P):
    rows = OrderedDict()
    for c in P.get('barh'):
        r = c.ctx['r']
        d = rows.setdefault(r.predictor, dict(drug=r.name, predictor=r.predictor,
                                              row_label=tick_at(c.ax, 'y', c.ctx['k']), per_drug_correlation=r.raw,
                                              r_pg=r.r_pg, r_zg=r.r_zg))
        d['selective_term' if 'left' in c.kw else 'axis_term'] = num(c.args[1])[0]
    cols = OrderedDict(drug='example drug', predictor='prediction of the drug (row)', row_label='row label',
                       per_drug_correlation='per-drug correlation r_pz of the prediction with the profile (number at '
                                            'right; 1 for the perfect predictor)',
                       r_pg="correlation of the prediction with the axis g", r_zg='correlation of the profile with g',
                       axis_term='axis term r_pg·r_zg (grey bar)',
                       selective_term='selective term r_pz - r_pg·r_zg (orange bar, stacked on the axis term)')
    df = T(list(rows.values()), cols)
    k = [('predictor', 'predictor')]
    chk = [check(df, c, f'{EX1}/terms.csv', k, u) for c, u in [('axis_term', 'axis_term'),
                                                              ('selective_term', 'selective_term'),
                                                              ('per_drug_correlation', 'raw')]]
    return df, cols, chk, ''


@conv('panel_rank_shift')
def c_rank_shift(P):
    leg = P.legend
    labelled = {c.ctx['m'] for c in P.get('text') if 'm' in c.ctx}
    rows = []
    for c in P.get('plot'):
        m = c.ctx['m']
        y = num(c.args[1])
        rows.append(dict(model=m, model_label=plain(M.NAMES.get(m, m)),
                         group=legend_label(leg, col=c.kw.get('color')) or 'other released model',
                         rank_raw=y[0], rank_selective=y[1], labelled_in_figure=m in labelled))
    cols = OrderedDict(model='released DrEval model', model_label='name in the figure',
                       group='colour: drug-agnostic (purple), EnsembleMF (blue), other released model (grey)',
                       rank_raw='rank by median per-drug (raw) correlation (left)',
                       rank_selective='rank by median selective score (right)',
                       labelled_in_figure='name and selective rank printed in the panel')
    df = T(rows, cols)
    R = 'results/revision_round2/rankings/rank_intervals.csv'
    chk = [check(df, 'rank_raw', R, [('model', 'model')], 'rank', up_filter=lambda U: U[U.metric == 'raw']),
           check(df, 'rank_selective', R, [('model', 'model')], 'rank', up_filter=lambda U: U[U.metric == 'selective'])]
    return df, cols, chk, 'Rank intervals: Supplementary Table (rank_lo, rank_hi in the upstream file).'


# ---------------------------------------------------------------------------------------------- Figure 5
TPD = 'results/revision_round4/transfer_prediction'
TPB = rel(M.TPB)   # pooled correlations as the figure reads them
SIMD = 'results/revision_round4/simulation'
X_NAME = {'raw': ('raw_score_in_CTRPv2', 'raw leaderboard score in CTRP v2 (median per-drug correlation; x)'),
          'S': ('selective_capture_in_CTRPv2', 'selective capture in CTRP v2 (x)')}


@conv('panel_transfer_scatter')
def c_transfer_scatter(P):
    rows, xname, r, xcol = [], None, None, None
    for c in P.get('scatter'):
        xcol = c.ctx['xcol']
        xname = X_NAME.get(xcol, (f'{xcol}_in_CTRPv2', f'{xcol} within CTRP v2 (x)'))
        r = c.ctx['r']
        for m, x, y in zip(sidx(c.args[0]), num(c.args[0]), num(c.args[1])):
            rows.append({'model': m, 'drug_agnostic': bool(c.ctx['agn']), xname[0]: x, 'selective_both_other_datasets': y,
                         'element': 'model'})
    if r is not None:
        rows.append(dict(model='pooled Spearman correlation (panel text)', predictor=r.predictor, spearman_r=r.r_pooled,
                         spearman_ci_lo=r.lo, spearman_ci_hi=r.hi, targets=r.targets, element='summary'))
    cols = OrderedDict(model='released DrEval leaderboard model', drug_agnostic='drug-agnostic (purple) or learned (blue)')
    cols[xname[0]] = xname[1]
    cols.update(selective_both_other_datasets='selective-both the same predictions keep in the other datasets (mean '
                                              'over the 4 CurveCurator ln IC50 targets; y)',
                element='model, or the pooled Spearman correlation printed in the panel',
                predictor='within-CTRP v2 statistic correlated', spearman_r='pooled Spearman correlation over targets',
                spearman_ci_lo='95% interval (bootstrap over models, drugs and cell lines), lower',
                spearman_ci_hi='95% interval, upper', targets='targets pooled')
    df = T(rows, cols)
    A = upstream(f'{TPD}/fig_a_model_scatter.csv')
    A = A[A.benchmark == 'leaderboard'].groupby('model')
    mods = df[df.element == 'model']
    src = {'raw': 'raw_median'}.get(xcol, xcol)
    chk = [check_values('selective_both_other_datasets', mods.selective_both_other_datasets,
                        A.selB_c.mean().loc[mods.model], f'{TPD}/fig_a_model_scatter.csv',
                        'selB_c (mean over targets per model, leaderboard rows)',
                        note='recomputation: mean over targets as in the panel code'),
           check_values(xname[0], mods[xname[0]], A[src].first().loc[mods.model], f'{TPD}/fig_a_model_scatter.csv',
                        f'{src} (per model, leaderboard rows)'),
           check(df, 'spearman_r', f'{TPB}/fig_b_pooled_correlations.csv', [('predictor', 'predictor')], 'r_pooled',
                 mask=df.element == 'summary',
                 up_filter=lambda U: U[(U.group == 'leaderboard') & (U.model_set == 'all') & (U.outcome == 'selB_c')
                                       & (U.method == 'spearman') & (U.interval == 'full')])]
    return df, cols, chk, ''


@conv('panel_transfer_forest')
def c_transfer_forest(P):
    leg = P.legend
    rows = []
    for c in P.get('errorbar'):
        x, r = c.ctx, c.ctx['r']
        rows.append(dict(benchmark=plain(x['glab']), group=x['g'],
                         predictor=legend_label(leg, marker='o', col=c.kw.get('color')), predictor_code=x['pr'],
                         outcome=legend_label(leg, col=M.INK2, mfc=M.INK2 if x['filled'] else 'white'),
                         outcome_code=x['out'], spearman_r=num(c.args[0])[0], ci_lo=r.lo, ci_hi=r.hi, targets=r.targets))
    cols = OrderedDict(benchmark='benchmark (row group)', group='group in the upstream file',
                       predictor='within-dataset statistic (colour): raw score, DrEval normalized metric, selective capture',
                       predictor_code='predictor in the upstream file',
                       outcome='target selective-both (filled) or target raw accuracy (open)',
                       outcome_code='outcome in the upstream file',
                       spearman_r='pooled Spearman correlation across models (over targets)',
                       ci_lo='95% interval, lower', ci_hi='95% interval, upper', targets='targets pooled')
    df = T(rows, cols)
    chk = [check(df, 'spearman_r', f'{TPB}/fig_b_pooled_correlations.csv',
                 [('group', 'group'), ('predictor_code', 'predictor'), ('outcome_code', 'outcome')], 'r_pooled',
                 up_filter=lambda U: U[(U.model_set == 'all') & (U.method == 'spearman') & (U.interval == 'full')])]
    return df, cols, chk, ''


@conv('panel_sim_equal')
def c_sim_equal(P):
    rows = []
    for c in P.get('bar'):
        x, r = c.ctx, c.ctx['r']
        rows.append(dict(predictor=x['lab'], predictor_code=x['p_'], statistic=x['slab'], stat_code=x['st'],
                         D=r.D, N=r.N, replicates=r.replicates, mean_estimate=num(c.args[1])[0]))
    cols = OrderedDict(predictor='simulated predictor with the same raw score (x group)',
                       predictor_code='predictor in the upstream file',
                       statistic='raw score (grey), axis capture (light blue), selective capture (blue)',
                       stat_code='stat in the upstream file', D='simulated drugs', N='simulated cell lines',
                       replicates='simulation replicates', mean_estimate='estimate, mean over replicates (bar)')
    df = T(rows, cols)
    chk = [check(df, 'mean_estimate', f'{SIMD}/fig_1_2_recovery_coverage.csv',
                 [('D', 'D'), ('predictor_code', 'predictor'), ('stat_code', 'stat')], 'mean_estimate')]
    return df, cols, chk, ''


@conv('panel_sim_recovery')
def c_sim_recovery(P):
    leg = P.legend
    rows = []
    for c in P.get('scatter'):
        d = c.ctx['d']
        for i, x, y in zip(sidx(c.args[0]), num(c.args[0]), num(c.args[1])):
            r = d.loc[i]
            rows.append(dict(statistic=legend_label(leg, col=c.kw.get('color')), stat_code=c.ctx['st'],
                             panel_size=legend_label(leg, marker=c.kw.get('marker', 'o'), col=M.INK2), D=r.D, N=r.N,
                             predictor=r.predictor, true_capture=x, estimated_capture=y,
                             coverage_95=r.coverage_design))
    cols = OrderedDict(statistic='axis capture (light blue) or selective capture (blue)', stat_code='stat in the upstream file',
                       panel_size='simulated panel (circle: 150 drugs × 500 cells; triangle: 30 × 100)',
                       D='simulated drugs', N='simulated cell lines', predictor='simulated predictor',
                       true_capture='true capture of the design (x)', estimated_capture='estimate, mean over 300 '
                                                                                        'replicates (y)',
                       coverage_95='coverage of the 95% interval (range printed in the panel)')
    df = T(rows, cols)
    k = [('D', 'D'), ('predictor', 'predictor'), ('stat_code', 'stat')]
    U = f'{SIMD}/fig_1_2_recovery_coverage.csv'
    chk = [check(df, 'estimated_capture', U, k, 'mean_estimate'), check(df, 'true_capture', U, k, 'truth_design')]
    return df, cols, chk, 'Dashed line: identity.'


@conv('panel_sim_leakage')
def c_sim_leakage(P):
    rows = OrderedDict()
    for c in P.get(['scatter', 'plot']):
        if 'pr' not in c.ctx or 'd' not in c.ctx:
            continue
        d = c.ctx['d']
        for i, x, y in zip(sidx(c.args[0]), num(c.args[0]), num(c.args[1])):
            row = rows.setdefault(i, dict(predictor=plain(c.ctx['lab']), predictor_code=c.ctx['pr'], D=d.loc[i, 'D'],
                                          axis_split_half=x))
            row['leakage' if c.method == 'scatter' else 'predicted_leakage'] = y
    for c in P.get('axvline'):
        rows['ref'] = dict(predictor="leaderboard axis (dotted line: median split-half reliability)",
                           axis_split_half=num(c.args[0])[0])
    cols = OrderedDict(predictor='simulated predictor (colour), or the reference line',
                       predictor_code='predictor in the upstream file', D='simulated drugs per panel',
                       axis_split_half='split-half reliability of the axis (x)',
                       leakage='leakage of the axis into selective capture (points)',
                       predicted_leakage='prediction A·m(1-ρ)/[(1-m)ρ] with the latent reliability (line)')
    df = T(list(rows.values()), cols)
    U = f'{SIMD}/fig_5_axis_reliability_leakage.csv'
    k = [('predictor_code', 'predictor'), ('D', 'D')]
    m = df.predictor_code.notna()
    X = upstream(f'{SIMD}/realdata_leakage_correction.csv')
    ref = df[~m]
    chk = [check(df, 'leakage', U, k, 'leakage', mask=m), check(df, 'predicted_leakage', U, k,
                                                                 'predicted_leakage_latent_rel', mask=m),
           check_values('axis_split_half (leaderboard line)', ref.axis_split_half,
                        [X[X.benchmark == 'leaderboard'].axis_split_half.median()] * len(ref),
                        f'{SIMD}/realdata_leakage_correction.csv', 'axis_split_half (median, leaderboard rows)',
                        note='recomputation: median as in the panel code')]
    return df, cols, chk, ''


# ---------------------------------------------------------------------------------------------- generic fallback
# ---------------------------------------------------------------------------------------------- round 5 panels
CAPP = 'results/revision_round2/capture/capture_pooled.csv'
DEPD = 'results/revision_round5/crispr_dependency'


@conv('panel_amplification')
def c_amplification(P):
    rows = []
    for c in P.get('errorbar'):
        r = c.ctx['row']
        rows.append(dict(benchmark=r['benchmark'], setting=r['setting'], domain=r['domain'],
                         group=r.get('group', ''), models=int(r['models']),
                         perfect_share=num(c.args[0])[0], perfect_share_lo=float(r['perfect_share_lo']),
                         perfect_share_hi=float(r['perfect_share_hi']), axis_share=num(c.args[1])[0],
                         axis_share_lo=float(r['axis_share_lo']), axis_share_hi=float(r['axis_share_hi']),
                         share_ratio=float(r['share_ratio'])))
    cols = OrderedDict(benchmark='benchmark (leaderboard, retrain, improve; crispr: gene dependency)',
                       setting='dataset and endpoint', domain='drug response or gene dependency',
                       group='model pool in the upstream file', models='models pooled',
                       perfect_share="perfect predictor's axis share (x; mean r_zg²)", perfect_share_lo='95% interval, lower',
                       perfect_share_hi='95% interval, upper', axis_share="pooled learned models' axis share (y)",
                       axis_share_lo='95% interval, lower', axis_share_hi='95% interval, upper',
                       share_ratio="models' axis share / perfect predictor's")
    df = T(rows, cols)
    drug = df.domain == 'drug response'
    keep = lambda U: U[~U.group.str.contains('ridges|minus|this study')]  # noqa: E731
    dep = lambda U: U[(U.config == 'primary') & (U.group == 'learned')].assign(setting='DepMap_Chronos')  # noqa: E731
    chk = [check(df, 'perfect_share', CAPP, [('benchmark', 'benchmark'), ('setting', 'setting')], 'perfect_share',
                 mask=drug, up_filter=keep),
           check(df, 'axis_share', CAPP, [('benchmark', 'benchmark'), ('setting', 'setting')], 'axis_share',
                 mask=drug, up_filter=keep),
           check(df, 'perfect_share', f'{DEPD}/pooled.csv', [('setting', 'setting')], 'perfect_share', mask=~drug,
                 up_filter=dep),
           check(df, 'axis_share', f'{DEPD}/pooled.csv', [('setting', 'setting')], 'axis_share', mask=~drug,
                 up_filter=dep)]
    return df, cols, chk, ''


@conv('panel_dep_capture')
def c_dep_capture(P):
    rows, tick = [], {}
    for c in P.get('scatter', marker='_'):
        tick[c.ctx['row']['model']] = num(c.args[1])[0]
    for c in P.get('scatter'):
        if c.kw.get('marker') == '_':
            continue
        r = c.ctx['row']
        rows.append(dict(model=r['model'], kind=r['kind'], axis_capture=num(c.args[0])[0],
                         selective_capture=num(c.args[1])[0], selective_capture_lineage_adjusted=tick.get(r['model'])))
    cols = OrderedDict(model='model', kind='learned, gene-agnostic or tissue-by-gene baseline',
                       axis_capture="axis capture (x)", selective_capture='selective capture (y)',
                       selective_capture_lineage_adjusted='selective capture with lineage also partialled out (tick)')
    df = T(rows, cols)
    pri = lambda U: U[U.config == 'primary']  # noqa: E731
    chk = [check(df, c_, f'{DEPD}/models.csv', [('model', 'model')], u_, up_filter=pri)
           for c_, u_ in [('axis_capture', 'axis_capture'), ('selective_capture', 'selective_capture'),
                          ('selective_capture_lineage_adjusted', 'selective_capture_lin')]]
    return df, cols, chk, ''


@conv('panel_dep_predictability')
def c_dep_predictability(P):
    bars = P.get('barh')
    Pr = bars[0].ctx['Pr']
    rows = [dict(model=m, genes_r_ge_0_3=num(bars[0].args[1])[i], of_which_axis_term_ge_half_r=num(bars[1].args[1])[i],
                 lo=float(Pr.n_pass_raw_lo.iloc[i]), hi=float(Pr.n_pass_raw_hi.iloc[i]))
            for i, m in enumerate(Pr.model)]
    cols = OrderedDict(model='model', genes_r_ge_0_3='genes with out-of-fold per-gene r >= 0.3 (light bar)',
                       of_which_axis_term_ge_half_r='of which the axis term is at least half of r (dark bar)',
                       lo='95% interval, lower', hi='95% interval, upper')
    df = T(rows, cols)
    f = lambda U: U[(U.config == 'primary') & (U.threshold == 0.3)]  # noqa: E731
    chk = [check(df, 'genes_r_ge_0_3', f'{DEPD}/predictability.csv', [('model', 'model')], 'n_pass_raw', up_filter=f),
           check(df, 'of_which_axis_term_ge_half_r', f'{DEPD}/predictability.csv', [('model', 'model')],
                 'n_axis_term_ge_half_r', up_filter=f)]
    return df, cols, chk, ''


@conv('panel_dep_axis')
def c_dep_axis(P):
    c = P.get('errorbar')[0]
    keep, v = c.ctx['keep'], c.ctx['v']
    rows = [dict(covariate=k, label=lab, pearson_r=num(c.args[0])[i], lo=float(v.lo.iloc[i]), hi=float(v.hi.iloc[i]))
            for i, (k, lab) in enumerate(keep)]
    cols = OrderedDict(covariate='covariate in the upstream file', label='row label',
                       pearson_r='Pearson r with the active-gene dependency axis', lo='95% interval, lower',
                       hi='95% interval, upper')
    df = T(rows, cols)
    f = lambda U: U[(U.axis == 'active_axis') & (U.statistic == 'pearson_r')]  # noqa: E731
    chk = [check(df, 'pearson_r', f'{DEPD}/axis.csv', [('covariate', 'with')], 'value', up_filter=f)]
    return df, cols, chk, ''


def c_generic(P):
    rows = []
    for c in P.calls:
        if c.method not in DATA_METHODS:
            continue
        base = dict(drawn_by=c.pfunc, method=c.method, label=plain(c.kw.get('label', '')),
                    color=color(c.kw.get('color', c.kw.get('facecolor', ''))) if c.kw.get('color', c.kw.get('facecolor')) is not None else '')
        for k, v in c.ctx.items():
            if isinstance(v, (str, int, float, bool, np.integer, np.floating)) and not k.startswith('_'):
                base['ctx_' + k] = v
        a = c.args
        if c.method in ('axvline', 'axhline'):
            rows.append(dict(base, value=num(a[0] if a else c.kw.get('x', c.kw.get('y', 0)))[0]))
            continue
        if len(a) < 2:
            continue
        x, y = num(a[0]), num(a[1])
        if len(x) != len(y):
            continue
        xlo, xhi = err_bounds(x, c.kw.get('xerr'))
        ylo, yhi = err_bounds(y, c.kw.get('yerr'))
        for j in range(len(x)):
            rows.append(dict(base, x=x[j], y=y[j], x_lo=xlo[j], x_hi=xhi[j], y_lo=ylo[j], y_hi=yhi[j],
                             x_tick=tick_at(c.ax, 'x', x[j]), y_tick=tick_at(c.ax, 'y', y[j])))
    cols = OrderedDict(drawn_by='drawing function', method='matplotlib call', label='legend label', color='colour',
                       x='x', y='y', x_lo='x interval, lower', x_hi='x interval, upper', y_lo='y interval, lower',
                       y_hi='y interval, upper', x_tick='nearest x tick label', y_tick='nearest y tick label',
                       value='position of a reference line')
    return T(rows, cols), cols, [], 'Generic export (no panel-specific converter): one row per plotted point.'


# =============================================================================================== captions and figures
def tex_macros():
    s = READ_TEXT(MANUSCRIPT / 'supp_refs.tex')
    return dict(re.findall(r'\\newcommand\{\\([A-Za-z]+)\}\{([^{}]*)\}', s))


def _braced(s, i):
    depth = 0
    for j in range(i, len(s)):
        if s[j] == '{':
            depth += 1
        elif s[j] == '}':
            depth -= 1
            if depth == 0:
                return s[i + 1:j]
    raise ValueError('unbalanced braces')


def figure_blocks(path):
    """[(file stem, caption, label)] of the figure environments with a graphic and a caption, in order."""
    s = re.sub(r'(?<!\\)%.*', '', READ_TEXT(path))
    out = []
    for m in re.finditer(r'\\begin\{figure\}(.*?)\\end\{figure\}', s, re.S):
        b = m.group(1)
        gm = re.search(r'\\includegraphics(?:\[[^\]]*\])?\{([^}]*)\}', b)
        i = b.find('\\caption{')
        if not gm or i < 0:
            continue                 # e.g. Box 1 (no graphic, no caption, not numbered)
        lm = re.search(r'\\label\{([^}]*)\}', b)
        out.append(dict(stem=Path(gm.group(1)).stem, caption=_braced(b, i + len('\\caption')),
                        label=lm.group(1) if lm else ''))
    return out


MARK_IN = re.compile(r'\\textbf\{([a-h](?:(?:--|,)[a-h])*),\}\s*')     # \textbf{a,}  \textbf{a--c,}  \textbf{a,b,}
MARK_OUT = re.compile(r'\\textbf\{([a-h](?:--[a-h])?)\},\s*')            # \textbf{a},  (supplementary captions)


def _letters(spec):
    out = []
    for part in spec.split(','):
        a, _, b = part.partition('--')
        out += [chr(c) for c in range(ord(a), ord(b or a) + 1)]
    return out


def caption_parts(cap, macros):
    cap = re.sub(r'\\((?:st|sf)[A-Za-z]+)', lambda m: macros.get(m.group(1), m.group(0)), cap)
    marks = list(MARK_IN.finditer(cap)) or list(MARK_OUT.finditer(cap))
    head = cap[:marks[0].start()] if marks else cap
    tm = re.match(r'\s*\\textbf\{', head)
    title = _braced(head, head.find('{')) if tm else ''
    rest = head[head.find(title) + len(title) + 1:] if title else head
    parts = {}
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(cap)
        for L in _letters(m.group(1)):
            parts.setdefault(L, []).append(cap[m.end():end].strip())
    return plain(title), plain(rest), {L: plain(' '.join(v)) for L, v in parts.items()}


def drawing_function(stem):
    """The function of 64_figures_revision.py (preferred) or 13_figures.py that saves this figure file."""
    for mod in (M, F):
        for name, fn in vars(mod).items():
            if name.startswith('panel_') or not inspect.isfunction(fn) or fn.__module__ != mod.__name__:
                continue
            try:
                src = inspect.getsource(fn)
            except (OSError, TypeError):
                continue
            if re.search(r"save\(fig,\s*'%s'\)" % re.escape(stem), src):
                return mod, name
    return None, None


def figure_list():
    macros = tex_macros()
    figs = []
    for n, b in enumerate(figure_blocks(MANUSCRIPT / 'sections' / 'figures.tex'), 1):
        figs.append(dict(kind='main', number=n, key=f'Fig{n}', sheet_prefix=f'Fig{n}', book=f'Source_Data_Fig{n}',
                         display=f'Figure {n}', **b))
    for n, b in enumerate(figure_blocks(MANUSCRIPT / 'sections' / 'extended_data.tex'), 1):
        figs.append(dict(kind='ed', number=n, key=f'EDFig{n}', sheet_prefix=f'EDFig{n}',
                         book=f'Source_Data_ExtendedDataFig{n}', display=f'Extended Data Figure {n}', **b))
    sb = figure_blocks(MANUSCRIPT / 'supplement_figures.tex')
    for k, b in enumerate(sb, 1):
        lab = b['label'].split(':')[-1]
        n = int(macros.get('sf' + lab, k))
        figs.append(dict(kind='supp', number=n, key=f'SFig{n}', sheet_prefix=f'SFig{n}', book=f'Source_Data_SuppFig{n}',
                         display=f'Supplementary Figure {n}', **b))
    for f in figs:
        f['title'], f['head'], f['parts'] = caption_parts(f['caption'], macros)
        f['module'], f['function'] = drawing_function(f['stem'])
    return figs


# =============================================================================================== panels, reads
READ_OVERRIDE_FUNCS = {'target_rows': 'de'}   # Fig. 2 figure-level helper whose rows are drawn in panels d and e
READ_OVERRIDE_FILES = {                        # files read at the top of single-function supplementary figures
    ('supp_figure3_dreval', 'per_drug.csv.gz'): 'a', ('supp_figure3_dreval', 'model_summary.csv'): 'bc',
    ('supp_figure5_value_of_information', 'l1000_two_way.csv'): 'a',
    ('supp_figure5_value_of_information', 'retrieval_baselines.csv'): 'ab',
    ('supp_figure5_value_of_information', 'l1000_incremental.csv'): 'c',
    ('supp_figure_worked_example', 'resid_CTRP.parquet'): 'b', ('supp_figure_worked_example', '*'): 'ab'}


def build_panels(run):
    by = OrderedDict()
    for c in run.calls:
        by.setdefault(run.panel_of(c.ax), []).append(c)
    out = OrderedDict()
    for L in sorted(by):
        if any(c.method in DATA_METHODS for c in by[L]):
            out[L] = Panel(run, L, by[L])
    return out


def panel_reads(run, panels):
    tok = {}
    for L, P in panels.items():
        for c in P.calls:
            if c.token:
                tok.setdefault(c.token, set()).add(L)
    out = {L: OrderedDict() for L in panels}
    for r in run.reads:
        Ls = tok.get(r['token']) if r['token'] else None
        if Ls is None:
            hit = [READ_OVERRIDE_FUNCS[f] for f in r['stack'] if f in READ_OVERRIDE_FUNCS]
            if hit:
                Ls = set(hit[0])
            elif (run.stem, r['path'].name) in READ_OVERRIDE_FILES:
                Ls = set(READ_OVERRIDE_FILES[(run.stem, r['path'].name)])
            elif (run.stem, '*') in READ_OVERRIDE_FILES:
                Ls = set(READ_OVERRIDE_FILES[(run.stem, '*')])
            elif r['last_label'] is not None:
                Ls = {r['last_label']}
            else:
                Ls = set(panels)
        for L in Ls:
            if L in out:
                via = ' <- '.join(f + '()' for f in r['stack'][:3])
                vias = out[L].setdefault(rel(r['path']), [])
                if via not in vias:
                    vias.append(via)
    return out


def convert(P):
    fn = CONVERTERS.get(f'{P.func}:{P.letter}') or CONVERTERS.get(P.func) or c_generic
    try:
        df, cols, checks, note = fn(P)
    except Exception as e:      # the panel code changed under the converter: keep exporting, flag it
        df, cols, checks, note = c_generic(P)
        checks = [dict(check='converter', exported_column='', upstream_file='', upstream_column='', key='',
                       n_checked=0, n_exact_match=0, n_unmatched_key=0, max_abs_diff=np.nan, result='FAIL',
                       note=f'{fn.__name__}() failed ({type(e).__name__}: {e}); generic export written instead')]
        fn = c_generic
    return df.reset_index(drop=True), cols, checks, note, fn.__name__


def categorical_axes(P):
    """Axes whose x or y tick labels were set by the panel code (positions there are categories, not data)."""
    cat = set()
    for c in P.run.calls:
        if c.method in ('set_xticklabels', 'set_yticklabels'):     # including [] (labels shown on a sibling axes)
            cat.add((id(c.ax), c.method[4]))
    return cat


def coverage(P, df):
    """Share of the values drawn as data marks (non-categorical coordinates) that appear in the exported table."""
    cat = categorical_axes(P)
    drawn = []
    for c in P.calls:
        a = c.args
        if c.method in ('scatter', 'errorbar') and len(a) >= 2:
            if (id(c.ax), 'x') not in cat:
                drawn += list(num(a[0]))
            if (id(c.ax), 'y') not in cat:
                drawn += list(num(a[1]))
        elif c.method == 'plot' and len(a) >= 2 and ('marker' in c.kw or (len(a) > 2 and isinstance(a[2], str)
                                                                            and re.search(r'[o.sD^v<>x+*ph]', a[2]))):
            if (id(c.ax), 'x') not in cat:          # lines drawn with markers are data series, not reference lines
                drawn += list(num(a[0]))
            if (id(c.ax), 'y') not in cat:
                drawn += list(num(a[1]))
        elif c.method in ('bar', 'barh') and len(a) >= 2:
            drawn += list(num(a[1]))
        elif c.method == 'boxplot':
            drawn += [v for d in a[0] for v in num(d)]
        elif c.method == 'imshow':
            drawn += list(np.asarray(a[0], float).ravel())
    drawn = [v for v in drawn if np.isfinite(v)]
    have = set()
    for col in df.columns:
        if pd.api.types.is_numeric_dtype(df[col]) and not pd.api.types.is_bool_dtype(df[col]):
            have |= set(df[col].dropna().astype(float))
    n = len(drawn)
    ok = sum(v in have for v in drawn)
    return dict(check='coverage of drawn values', exported_column='all', upstream_file='(figure code)',
                upstream_column='', key='', n_checked=n, n_exact_match=ok, n_unmatched_key=0, max_abs_diff=np.nan,
                result='not applicable' if n == 0 else ('PASS' if ok == n else 'FAIL'),
                note='every value drawn by scatter/errorbar/bar/barh/boxplot/imshow (and lines with markers) on a '
                     'non-categorical axis is in the exported table (exact)' if n else
                     'no scatter/errorbar/bar/boxplot/imshow marks in this panel (lines only)')


# =============================================================================================== outputs
def sha256(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def snapshot(d):
    return {p.name: (p.stat().st_size, p.stat().st_mtime_ns, sha256(p)) for p in sorted(Path(d).iterdir()) if p.is_file()}


KNOWN_PNG_DIFFS = {   # differences inspected by eye; the note is used only while the difference signature is unchanged
    'figure5_transfer_simulation': ((3470, (231, 526, 222, 598)),
                                    'inspected 2026-10-15: only the position of the panel a legend differs (the code '
                                    'moved it after the published file was rendered); all data marks identical. '
                                    'Re-running 64_figures_revision.py 5 refreshes the published file')}


def png_compare(png, stem):
    pub = FIGDIR / f'{stem}.png'
    row = dict(check='regenerated figure vs published PNG', exported_column='', upstream_file=rel(pub),
               upstream_column='', key='', n_checked=1, n_unmatched_key=0, max_abs_diff=np.nan)
    if not pub.exists():
        return dict(row, n_exact_match=0, result='NO PUBLISHED FILE', note='')
    if hashlib.sha256(png).hexdigest() == sha256(pub):
        return dict(row, n_exact_match=1, result='PASS', max_abs_diff=0.0,
                    note='byte-identical: the exported values are those of the published figure')
    a, b = mpimg.imread(io.BytesIO(png)), mpimg.imread(pub)
    if a.shape != b.shape:
        return dict(row, n_exact_match=0, result='DIFFERS', note=f'image size {a.shape[:2]} vs published {b.shape[:2]}: '
                                                                    'the published figure predates the current inputs or code')
    d = np.abs(a - b).sum(-1) > 0
    ys, xs = np.where(d)
    sig = (int(d.sum()), (int(xs.min()), int(xs.max()), int(ys.min()), int(ys.max())))
    known = KNOWN_PNG_DIFFS.get(stem)
    extra = known[1] if known and known[0] == sig else ('published file older than the drawing code or inputs; '
                                                        're-run the figure script to refresh it')
    return dict(row, n_exact_match=0, result='DIFFERS',
                note=f'{sig[0]} of {d.size} pixels differ, within x {sig[1][0]}-{sig[1][1]}, y {sig[1][2]}-{sig[1][3]}: '
                     + extra)


def func_file(name):
    for mod in (M, F):
        fn = getattr(mod, name, None)
        fn = getattr(fn, '__wrapped__', fn)
        if inspect.isfunction(fn) and fn.__module__ == mod.__name__:
            return f'scripts/{Path(mod.__file__).name}'
    return ''


def readme_rows(fig, info, png_check):
    mod_file = f"scripts/{Path(fig['module'].__file__).name}" if fig['module'] else ''
    rows = [dict(sheet='', field='figure', name=fig['display'], description=fig['title']),
            dict(sheet='', field='published file', name=f"figures/{fig['stem']}.pdf (.png)", description=''),
            dict(sheet='', field='drawn by', name=f"{mod_file}: {fig['function']}()", description=''),
            dict(sheet='', field='generated', name=datetime.date.today().isoformat(),
                 description='by scripts/99i_source_data.py from the figure code itself (matplotlib calls recorded while '
                             'the figure is drawn); re-run "cd scripts && ../../.venv/bin/python 99i_source_data.py" '
                             'after the figure or its inputs change'),
            dict(sheet='', field='published figure check', name=png_check['result'], description=png_check['note']),
            dict(sheet='', field='values', name='', description='Plotted values are exported exactly as drawn (full '
                 'double precision); interval bounds are taken from the source row the panel drew them from. Rows '
                 'follow the drawing order. A blank cell means not applicable to that row (e.g. no interval for a '
                 'reference line). Upstream paths are relative to the selectivity_ceiling project directory.')]
    if fig['head']:
        rows.append(dict(sheet='', field='caption (whole figure)', name='', description=fig['head']))
    for sheet, d in info.items():
        rows.append(dict(sheet=sheet, field='caption', name=f"panel {d['letter']}" if d['letter'] else 'figure',
                         description=fig['parts'].get(d['letter'], fig['head'] or fig['title'])))
        rows.append(dict(sheet=sheet, field='drawn by', name=f"{func_file(d['func'])}: {d['func']}()",
                         description=f"converter {d['converter']}() of scripts/99i_source_data.py"))
        for p, vias in d['reads'].items():
            rows.append(dict(sheet=sheet, field='upstream file', name=p, description='read by ' + '; '.join(vias)))
        if not d['reads']:
            rows.append(dict(sheet=sheet, field='upstream file', name='(none)', description='no file read by this panel'))
        if d['note']:
            rows.append(dict(sheet=sheet, field='note', name='', description=d['note']))
        rows.append(dict(sheet=sheet, field='rows', name=str(len(d['df'])), description=''))
        for c in d['df'].columns:
            rows.append(dict(sheet=sheet, field='column', name=c, description=d['cols'].get(c, '')))
    return pd.DataFrame(rows, columns=['sheet', 'field', 'name', 'description'])


def _exact_xlsx_numbers():
    """openpyxl writes numbers as '%.16g' (can drop the 17th significant digit); write the shortest repr that
    round-trips instead, so the workbook holds the exact doubles (only within this process)."""
    import math
    import openpyxl.cell._writer as w
    if getattr(w.safe_string, '_exact', False):
        return
    orig = w.safe_string

    def safe_string(v):
        if isinstance(v, float) and math.isfinite(v):
            return repr(float(v))
        return orig(v)
    safe_string._exact = True
    w.safe_string = safe_string


def write_book(path, readme, sheets):
    from openpyxl.styles import Alignment, Font
    _exact_xlsx_numbers()
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        readme.to_excel(xw, sheet_name='README', index=False)
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = 'A2'
            for cell in ws[1]:
                cell.font = Font(bold=True)
            if ws.title == 'README':
                for col, w in zip('ABCD', (9, 22, 48, 120)):
                    ws.column_dimensions[col].width = w
                for row in ws.iter_rows(min_row=2):
                    row[3].alignment = Alignment(wrap_text=True, vertical='top')
                    for c in row[:3]:
                        c.alignment = Alignment(vertical='top')
            else:
                for k, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), 1):
                    w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                    ws.column_dimensions[col[0].column_letter].width = min(max(10, w + 2), 50)


def _norm(v):
    if v is None or (isinstance(v, float) and np.isnan(v)) or (isinstance(v, str) and v == ''):
        return np.nan                       # empty text is read back as missing
    if isinstance(v, (bool, np.bool_)):
        return float(v)
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    return str(v)


def frames_equal(a, b):
    """Exact equality of two tables (numbers as doubles, booleans as 0/1, missing as NaN)."""
    if list(a.columns) != list(b.columns) or a.shape != b.shape:
        return False
    for c in a.columns:
        for x, y in zip(a[c].map(_norm), b[c].map(_norm)):
            if isinstance(x, float) and isinstance(y, float):
                if not (x == y or (np.isnan(x) and np.isnan(y))):
                    return False
            elif x != y:
                return False
    return True


def main():
    before = snapshot(FIGDIR)
    install()
    figs = figure_list()
    (OUT / 'csv').mkdir(parents=True, exist_ok=True)
    for p in list(OUT.glob('Source_Data_*.xlsx')) + list((OUT / 'csv').glob('*.csv')):
        p.unlink()
    validation, manifest, all_readme, missing = [], [], [], []
    for fig in figs:
        if fig['module'] is None:
            missing.append(fig)
            print('no drawing code found for', fig['stem'])
            continue
        run = run_figure(fig['key'], fig['stem'], fig['number'], getattr(fig['module'], fig['function']))
        png = [s for s in run.saved if s[0] == fig['stem']]
        pchk = png_compare(png[0][2], fig['stem']) if png else dict(result='NOT SAVED', note='')
        validation.append(dict(figure=fig['display'], sheet='', **pchk))
        panels = build_panels(run)
        reads = panel_reads(run, panels)
        info, sheets = OrderedDict(), OrderedDict()
        for L, P in panels.items():
            sheet = f"{fig['sheet_prefix']}{L}"
            df, cols, checks, note, cname = convert(P)
            info[sheet] = dict(letter=L, func=P.func, converter=cname, df=df, cols=cols, note=note, reads=reads[L])
            sheets[sheet] = df
            for ch in checks + [coverage(P, df)]:
                validation.append(dict(figure=fig['display'], sheet=sheet, **ch))
            if not checks:
                validation.append(dict(figure=fig['display'], sheet=sheet, check='value check', result='not applicable',
                                       note=note or 'no upstream file'))
        readme = readme_rows(fig, info, pchk)
        book = OUT / f"{fig['book']}.xlsx"
        write_book(book, readme, sheets)
        rd = OUT / 'csv' / f"{fig['book']}_README.csv"
        readme.to_csv(rd, index=False)
        manifest.append(dict(file=rel(rd).replace('manuscript/source_data/', ''), panel=f"{fig['sheet_prefix']} README",
                             rows=len(readme), sha256=sha256(rd)))
        back = pd.read_excel(book, sheet_name=None, engine='openpyxl')
        for sheet, df in sheets.items():
            p = OUT / 'csv' / f'{sheet}.csv'
            df.to_csv(p, index=False)
            manifest.append(dict(file=rel(p).replace('manuscript/source_data/', ''), panel=sheet, rows=len(df),
                                 sha256=sha256(p)))
            ok_csv = frames_equal(df, READ_CSV(p, float_precision='round_trip'))
            ok_xl = frames_equal(df, back[sheet])
            validation.append(dict(figure=fig['display'], sheet=sheet, check='written files read back',
                                   exported_column='all', upstream_file=f"csv/{sheet}.csv; {book.name}[{sheet}]",
                                   n_checked=2, n_exact_match=int(ok_csv) + int(ok_xl),
                                   result='PASS' if ok_csv and ok_xl else 'FAIL',
                                   note='CSV and xlsx sheet equal the exported table (exact)'))
        manifest.append(dict(file=book.name, panel=f"{fig['sheet_prefix']} (all sheets)",
                             rows=int(sum(len(d) for d in sheets.values())), sha256=sha256(book)))
        all_readme.append(readme.assign(workbook=book.name))
        print(f"{fig['display']}: " + ', '.join(f'{s} {len(d)}' for s, d in sheets.items()), flush=True)
    after = snapshot(FIGDIR)
    same = [k for k in before if after.get(k) == before[k]]
    validation.append(dict(figure='all', sheet='', check='published figures untouched', exported_column='',
                           upstream_file='figures/', upstream_column='size, mtime, sha256', key='',
                           n_checked=len(before), n_exact_match=len(same),
                           n_unmatched_key=len(set(after) ^ set(before)), max_abs_diff=np.nan,
                           result='PASS' if len(same) == len(before) and set(after) == set(before) else 'FAIL',
                           note='every file in figures/ has the same size, modification time and SHA-256 before and '
                                'after this run'))
    R = pd.concat(all_readme, ignore_index=True)
    R = R[['workbook', 'sheet', 'field', 'name', 'description']]
    R.to_csv(OUT / 'README.csv', index=False)
    V = pd.DataFrame(validation)
    vcols = ['figure', 'sheet', 'check', 'exported_column', 'upstream_file', 'upstream_column', 'key', 'n_checked',
             'n_exact_match', 'n_unmatched_key', 'max_abs_diff', 'result', 'note']
    V = V[[c for c in vcols if c in V.columns]]
    V.to_csv(OUT / 'validation.csv', index=False)
    for p in (OUT / 'README.csv', OUT / 'validation.csv'):
        manifest.append(dict(file=p.name, panel='all', rows=len(READ_CSV(p)), sha256=sha256(p)))
    pd.DataFrame(manifest, columns=['file', 'panel', 'rows', 'sha256']).to_csv(OUT / 'manifest.csv', index=False)
    print(V.result.value_counts().to_string())
    bad = V[~V.result.isin(['PASS'])]
    if len(bad):
        print(bad[['figure', 'sheet', 'check', 'exported_column', 'result', 'note']].to_string())
    for f in missing:
        print('NOT EXPORTED (no drawing code):', f['stem'])


def dump():
    install()
    which = [a for a in sys.argv[1:] if not a.startswith('--')]
    for fig in figure_list():
        if which and fig['key'] not in which:
            continue
        run = run_figure(fig['key'], fig['stem'], fig['number'], getattr(fig['module'], fig['function']))
        print('=' * 20, fig['key'], fig['stem'], fig['function'])
        for c in run.calls:
            if c.method in DATA_METHODS | {'text', 'annotate'}:
                ctx = {k: (v if isinstance(v, (str, int, float)) else type(v).__name__) for k, v in c.ctx.items()}
                print(run.panel_of(c.ax), c.pfunc, c.lineno, c.method, str(ctx)[:200])


if __name__ == '__main__':
    if '--dump' in sys.argv:
        dump()
    else:
        main()
